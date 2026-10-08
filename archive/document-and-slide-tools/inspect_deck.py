from pathlib import Path
from pptx import Presentation

deck = Path(r"D:\projects\car-detection-yolo\thesis\ارائه\ارائه_دفاع_نهایی_جلسه.pptx")
prs = Presentation(deck)
print(f"slides={len(prs.slides)} size={prs.slide_width}x{prs.slide_height}")
for i, slide in enumerate(prs.slides, 1):
    print(f"\n--- SLIDE {i} ---")
    for shape in sorted(slide.shapes, key=lambda s: (s.top, s.left)):
        if getattr(shape, "has_text_frame", False):
            text = " | ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if text:
                print(text)
        elif shape.shape_type == 13:
            print(f"[IMAGE {shape.width}x{shape.height}]")
        elif getattr(shape, "has_table", False):
            rows = []
            for row in shape.table.rows:
                rows.append(" || ".join(c.text.strip().replace("\n", " / ") for c in row.cells))
            print("[TABLE] " + " // ".join(rows))
