# -*- coding: utf-8 -*-
"""Slides 1 to 14."""
import os

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from deck import (A, AMBER, BLUE, CRIMSON, CW, EN, F1, F2, F3, FA, FS, GREEN,
                  H, INK, MARGIN, NAVY, ROOT, SLATE, TINT_AMBER, TINT_BLUE,
                  TINT_GREEN, TINT_GREY, TINT_RED, TOP, W, WHITE, SLIDES,
                  animate, box, fa_num, line, picture, prs, slide, table,
                  textbox, tile, _run)


def reg(fn):
    SLIDES.append(fn)
    return fn


def divider(num, title, sub, tint):
    sl = slide(bg=tint)
    tb, tf = textbox(sl, MARGIN, Inches(2.55), CW, Inches(2.4))
    line(tf, "بخش " + fa_num(num), 26, True, BLUE, first=True,
         align=PP_ALIGN.CENTER, after=14)
    line(tf, title, 46, True, NAVY, align=PP_ALIGN.CENTER, after=12)
    line(tf, sub, 24, False, SLATE, align=PP_ALIGN.CENTER, after=0)
    return sl


def white_divider(num, title, sub):
    """Section divider consistent with the first academic section."""
    sl = slide(bg=WHITE)
    mark = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, W - Inches(0.42), 0,
                               Inches(0.42), H)
    mark.fill.solid(); mark.fill.fore_color.rgb = NAVY
    mark.line.fill.background(); mark.shadow.inherit = False
    tb, tf = textbox(sl, Inches(1.25), Inches(1.75), W - Inches(2.5), Inches(0.65))
    line(tf, "بخش " + fa_num(num), 27, True, BLUE, first=True,
         align=PP_ALIGN.CENTER, after=0)
    tb, tf = textbox(sl, Inches(1.25), Inches(2.55), W - Inches(2.5), Inches(1.0))
    line(tf, title, 46, True, NAVY, first=True, align=PP_ALIGN.CENTER, after=0)
    rule = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.2), Inches(3.72),
                               Inches(4.93), Inches(0.035))
    rule.fill.solid(); rule.fill.fore_color.rgb = NAVY
    rule.line.fill.background(); rule.shadow.inherit = False
    tb, tf = textbox(sl, Inches(1.25), Inches(4.05), W - Inches(2.5), Inches(0.7))
    line(tf, sub, 24, False, SLATE, first=True, align=PP_ALIGN.CENTER, after=0)
    return sl


def academic_tile(sl, x, y, w, h, head, body, accent=NAVY, hs=25, bs=20):
    sh = box(sl, x, y, w, h, WHITE, accent)
    sh.line.width = Pt(1.7)
    line(sh.text_frame, head, hs, True, accent, first=True,
         align=PP_ALIGN.CENTER, after=4)
    if body:
        line(sh.text_frame, body, bs, False, SLATE,
             align=PP_ALIGN.CENTER, after=0)
    return sh


def plain_item(sl, x, y, w, h, head, body="", accent=NAVY,
               hs=24, bs=19, align=PP_ALIGN.RIGHT):
    """A strong text item without a decorative enclosing box."""
    marker = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 x + w - Inches(0.09), y + Inches(0.06),
                                 Inches(0.09), h - Inches(0.12))
    marker.fill.solid(); marker.fill.fore_color.rgb = accent
    marker.line.fill.background(); marker.shadow.inherit = False
    tb, tf = textbox(sl, x, y, w - Inches(0.28), h, MSO_ANCHOR.MIDDLE)
    line(tf, head, hs, True, accent, first=True, align=align,
         after=3 if body else 0)
    if body:
        line(tf, body, bs, False, SLATE, align=align, after=0)
    return tb


def strong_table(sl, rows, x, y, w, h, hs=20, bs=19, hcols=()):
    """High-contrast table for low-quality projectors."""
    shp = table(sl, rows, x, y, w, h, hs, bs, hcols=hcols)
    tbl = shp.table
    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            is_head = ri == 0
            is_label = (len(row.cells) - 1 - ci) in hcols and ri > 0
            cell.fill.solid()
            cell.fill.fore_color.rgb = (NAVY if is_head else
                                        (CRIMSON if is_label else
                                         (TINT_GREY if ri % 2 == 0 else WHITE)))
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = WHITE if (is_head or is_label) else INK
                    run.font.bold = is_head or is_label
            tc_pr = cell._tc.get_or_add_tcPr()
            for edge in ("lnL", "lnR", "lnT", "lnB"):
                old = tc_pr.find(A + edge)
                if old is not None:
                    tc_pr.remove(old)
                ln = etree.SubElement(tc_pr, A + edge)
                ln.set("w", "12700")
                solid = etree.SubElement(ln, A + "solidFill")
                colour = etree.SubElement(solid, A + "srgbClr")
                colour.set("val", "AAB5BE")
                etree.SubElement(ln, A + "prstDash").set("val", "solid")
    return shp


# ----------------------------------------------------------------------- 1
@reg
def cover():
    sl = slide(bg=WHITE)
    top_rule = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.16))
    top_rule.fill.solid(); top_rule.fill.fore_color.rgb = NAVY
    top_rule.line.fill.background(); top_rule.shadow.inherit = False

    logo = sl.shapes.add_picture(os.path.join(FS, "cover_image2.png"),
                                 Inches(0), Inches(0.40), height=Inches(1.05))
    logo.left = int((W - logo.width) / 2)

    tb, tf = textbox(sl, Inches(1.1), Inches(1.52), W - Inches(2.2), Inches(0.45))
    line(tf, "دانشکده مهندسی کامپیوتر", 21, False, SLATE, first=True,
         align=PP_ALIGN.CENTER, after=0)

    tb, tf = textbox(sl, Inches(1.0), Inches(2.05), W - Inches(2.0), Inches(1.75))
    line(tf, "دسته‌بندی اشیای ترافیکی با استفاده از YOLO", 34, True, NAVY,
         first=True, align=PP_ALIGN.CENTER, after=8)
    line(tf, "۶ دسته متفاوت: عابر، تریلر، اتوبوس، خودروی سبک، وانت و موتور",
         20, False, INK, align=PP_ALIGN.CENTER, after=8)
    line(tf, "پیاده‌سازی با معماری YOLOv11", 25, True, BLUE,
         align=PP_ALIGN.CENTER, after=0)

    rule = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.35), Inches(4.05),
                               Inches(4.63), Inches(0.035))
    rule.fill.solid(); rule.fill.fore_color.rgb = NAVY
    rule.line.fill.background(); rule.shadow.inherit = False

    rows = [("استاد راهنما", "دکتر رضا شمسایی"),
            ("استاد داور", "[نام استاد داور]"),
            ("دانشجو", "محمد امیر صادق‌زاده")]
    y = Inches(4.42)
    for lab, name in rows:
        tb, tf = textbox(sl, Inches(7.2), y, Inches(2.2), Inches(0.45))
        line(tf, lab, 20, False, SLATE, first=True, align=PP_ALIGN.RIGHT, after=0)
        tb, tf = textbox(sl, Inches(3.65), y, Inches(3.25), Inches(0.45))
        line(tf, name, 22, True, INK, first=True, align=PP_ALIGN.RIGHT, after=0)
        sep = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.65), y + Inches(0.47),
                                  Inches(5.75), Inches(0.012))
        sep.fill.solid(); sep.fill.fore_color.rgb = RGBColor(0xB8, 0xC0, 0xC7)
        sep.line.fill.background(); sep.shadow.inherit = False
        y += Inches(0.63)

# ---------------------------------------------------------------------- 2b
@reg
def scope():
    sl = slide("رده‌های مصوب و نهایی")
    rows = [["عنوان مصوب", "رده نهایی", "وضعیت"],
            ["عابر", "شخص", "گسترده‌تر"],
            ["اتوبوس", "اتوبوس", "یکسان"],
            ["خودروی سبک", "خودروی سبک", "یکسان"],
            ["موتور", "موتورسیکلت", "یکسان"],
            ["تریلر", "خودروی باری", "ادغام شد"],
            ["وانت", "خودروی باری", "ادغام شد"],
            ["—", "چراغ راهنمایی", "افزوده شد"]]
    strong_table(sl, rows, MARGIN, TOP, Inches(6.0), Inches(4.25),
                 19, 19, hcols=())

    items = [("محدودیت آغاز پروژه", "تصویب موضوع هم‌زمان با قطعی سراسری اینترنت"),
             ("انتخاب مجموعه‌داده", "انتخاب IADD به‌جای BDD100K با نظر استاد راهنما"),
             ("تلاش برای انطباق رده‌ها", "بازبرچسب‌گذاری کامل در زمان پروژه عملی نبود")]
    y = TOP + Inches(0.08)
    for head, body in items:
        plain_item(sl, Inches(6.95), y, Inches(5.75), Inches(1.06),
                   head, body, NAVY, 23, 19)
        y += Inches(1.33)
    plain_item(sl, MARGIN, Inches(5.70), CW, Inches(0.70),
               "مبنای نهایی: برچسب‌های موجود IADD", "", GREEN, 24, 19,
               align=PP_ALIGN.CENTER)

# ----------------------------------------------------------------------- 2
@reg
def agenda():
    sl = slide("فهرست ارائه")
    items = [(1, "مبانی و صورت مسئله"),
             (2, "روش پژوهش و آماده‌سازی داده"),
             (3, "نتایج و تحلیل"),
             (4, "جمع‌بندی و پیشنهادها")]
    y = TOP + Inches(0.16)
    for num, txt in items:
        sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, y, CW, Inches(1.08))
        sh.fill.solid(); sh.fill.fore_color.rgb = WHITE
        sh.line.color.rgb = NAVY; sh.line.width = Pt(1.6); sh.shadow.inherit = False
        nbox = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   W - MARGIN - Inches(0.92), y,
                                   Inches(0.92), Inches(1.08))
        nbox.fill.solid(); nbox.fill.fore_color.rgb = NAVY
        nbox.line.fill.background(); nbox.shadow.inherit = False
        line(nbox.text_frame, fa_num(num), 27, True, WHITE, first=True,
             align=PP_ALIGN.CENTER, after=0)
        tb, tf = textbox(sl, MARGIN + Inches(0.35), y + Inches(0.05),
                         CW - Inches(1.55), Inches(0.92), MSO_ANCHOR.MIDDLE)
        line(tf, txt, 28, True, NAVY, first=True, align=PP_ALIGN.RIGHT, after=0)
        y += Inches(1.25)

# ----------------------------------------------------------------------- 3
@reg
def d1():
    sl = slide(bg=WHITE)
    mark = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, W - Inches(0.42), 0,
                               Inches(0.42), H)
    mark.fill.solid(); mark.fill.fore_color.rgb = NAVY
    mark.line.fill.background(); mark.shadow.inherit = False
    tb, tf = textbox(sl, Inches(1.25), Inches(1.75), W - Inches(2.5), Inches(0.65))
    line(tf, "بخش ۱", 27, True, BLUE, first=True, align=PP_ALIGN.CENTER, after=0)
    tb, tf = textbox(sl, Inches(1.25), Inches(2.55), W - Inches(2.5), Inches(1.0))
    line(tf, "مبانی و صورت مسئله", 46, True, NAVY, first=True,
         align=PP_ALIGN.CENTER, after=0)
    rule = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.2), Inches(3.72),
                               Inches(4.93), Inches(0.035))
    rule.fill.solid(); rule.fill.fore_color.rgb = NAVY
    rule.line.fill.background(); rule.shadow.inherit = False
    tb, tf = textbox(sl, Inches(1.25), Inches(4.05), W - Inches(2.5), Inches(0.7))
    line(tf, "اهمیت مسئله، داده بومی و معماری مدل", 24, False, SLATE,
         first=True, align=PP_ALIGN.CENTER, after=0)

# ----------------------------------------------------------------------- 4
@reg
def problem():
    sl = slide("اهمیت تشخیص اشیای ترافیکی", size=28)
    items = [("مسئله", "تعیین رده و موقعیت هر شیء در تصویر"),
             ("کاربرد", "اطلاعات لازم برای سامانه‌های کمک‌راننده"),
             ("چالش محیط واقعی", "تغییر نور، فاصله، اندازه و هم‌پوشانی اشیا"),
             ("هدف پروژه", "تشخیص شش رده ترافیکی با ارزیابی قابل‌اتکا")]
    y = Inches(1.65)
    shapes = []
    for i, (head, body) in enumerate(items):
        x = Inches(6.95) if i % 2 == 0 else MARGIN
        yy = y if i < 2 else y + Inches(2.05)
        accent = GREEN if head == "هدف پروژه" else NAVY
        shapes.append(plain_item(sl, x, yy, Inches(5.75), Inches(1.45),
                                 head, body, accent, 26, 21))
    animate(sl, shapes)

# ----------------------------------------------------------------------- 5
@reg
def domain():
    sl = slide("تنظیم دقیق مدل با IADD", size=28)
    right = plain_item(sl, Inches(7.25), Inches(1.62), Inches(5.45), Inches(1.18),
                       "نقطه آغاز", "YOLOv11s با وزن‌های ازپیش‌آموخته روی COCO",
                       NAVY, 23, 21)
    tb, tf = textbox(sl, Inches(6.25), Inches(1.78), Inches(0.75), Inches(0.65),
                     MSO_ANCHOR.MIDDLE)
    line(tf, "←", 34, True, BLUE, first=True, align=PP_ALIGN.CENTER, after=0)
    left = plain_item(sl, MARGIN, Inches(1.62), Inches(5.45), Inches(1.18),
                      "مرحله پروژه", "تنظیم دقیق مدل با IADD",
                      GREEN, 23, 25)

    why = plain_item(sl, Inches(7.05), Inches(3.48), Inches(5.65), Inches(1.85),
                     "دلیل انتخاب IADD", "صحنه‌های واقعی ایران و تنوع شرایط جوی، زمانی و محیطی",
                     NAVY, 26, 21)
    goal = plain_item(sl, MARGIN, Inches(3.48), Inches(5.65), Inches(1.85),
                      "هدف تنظیم دقیق", "سازگاری مدل ازپیش‌آموخته با داده و شش رده هدف",
                      GREEN, 26, 21)
    animate(sl, [right, left, why, goal])

# ----------------------------------------------------------------------- 6
@reg
def iadd():
    sl = slide("داده‌های IADD در پروژه", size=27)
    rows = [["ویژگی", "IADD در مقاله", "نسخه برچسب‌دار در دسترس", "زیرمجموعه نهایی"],
            ["تصویر", "۹۷٬۵۲۸", "۸۳٬۰۴۲", "۳۷٬۱۴۲"],
            ["ویدیو", "۱۷۱", "۱۶۶", "۱۶۰"],
            ["رده", "شش", "شش", "شش"]]
    strong_table(sl, rows, MARGIN, TOP, CW, Inches(3.15), 20, 20, hcols=(0,))

    plain_item(sl, MARGIN, Inches(4.85), CW, Inches(0.72),
               "۱۴٬۴۸۶ تصویر بخش آزمون در نسخه دریافت‌شده فایل برچسب نداشتند",
               "", CRIMSON, 23, 19)

    plain_item(sl, MARGIN, Inches(5.62), Inches(5.75), Inches(0.96),
               "شش رده", "رده‌های اصلی IADD", NAVY, 22, 19)
    plain_item(sl, Inches(6.95), Inches(5.62), Inches(5.75), Inches(0.96),
               "هشت شهر ایران", "مبدأ گردآوری مجموعه اصلی", NAVY, 22, 19)

# ----------------------------------------------------------------------- 7
@reg
def arch():
    sl = slide("معماری YOLOv11")
    picture(sl, os.path.join(F1, "fig_1_2.png"), TOP - Inches(0.05),
            height=Inches(4.18))
    items = [("C3k2", "جایگزین C2f در ستون‌فقرات و گردن"),
             ("SPPF و C2PSA", "در مرز ستون‌فقرات و گردن"),
             ("سر بدون لنگر", "خروجی در سه مقیاس P3، P4 و P5")]
    w, gap = Inches(3.88), Inches(0.22)
    x = W - MARGIN - w
    shapes = []
    for head, body in items:
        shapes.append(plain_item(sl, x, Inches(5.78), w, Inches(0.74),
                                 head, body, NAVY, 21, 19,
                                 align=PP_ALIGN.CENTER))
        x -= (w + gap)
    animate(sl, shapes)


# ----------------------------------------------------------------------- 8
@reg
def metrics():
    sl = slide("معیارهای ارزیابی", size=26)
    items = [("دقت", "نسبت پیش‌بینی‌های درست"),
             ("فراخوانی", "نسبت اشیای واقعی یافت‌شده"),
             ("امتیاز F1", "میانگین هماهنگ دقت و فراخوانی"),
             ("mAP", "میانگین مساحت زیر منحنی دقت و فراخوانی")]
    w, gap = Inches(2.87), Inches(0.30)
    x = W - MARGIN - w
    for head, sub in items:
        tb, tf = textbox(sl, x, TOP + Inches(0.35), w, Inches(1.75), MSO_ANCHOR.MIDDLE)
        line(tf, head, 28, True, NAVY, first=True, align=PP_ALIGN.CENTER, after=8)
        line(tf, sub, 20, False, SLATE, align=PP_ALIGN.CENTER, after=0)
        x -= (w + gap)
    tb, tf = textbox(sl, Inches(6.95), Inches(4.45), Inches(5.75), Inches(1.25))
    line(tf, "mAP@0.5", 26, True, NAVY, first=True, after=6)
    line(tf, "آستانه هم‌پوشانی ۰٫۵", 21, False, SLATE, after=0)
    tb, tf = textbox(sl, MARGIN, Inches(4.45), Inches(5.75), Inches(1.25))
    line(tf, "mAP@0.5:0.95", 26, True, NAVY, first=True, after=6)
    line(tf, "میانگین ده آستانه از ۰٫۵ تا ۰٫۹۵", 21, False, SLATE, after=0)


# ----------------------------------------------------------------------- 9
@reg
def d2():
    white_divider(2, "روش پژوهش و آماده‌سازی داده",
                  "چالش‌های داده، ساخت زیرمجموعه و پیکربندی آموزش")


# ---------------------------------------------------------------------- 10
@reg
def arc():
    sl = slide("مسیر تکامل پروژه")
    picture(sl, os.path.join(FS, "slide_runs.png"), TOP - Inches(0.05),
            height=Inches(4.15))
    a, af = textbox(sl, Inches(8.85), Inches(5.72), Inches(3.85), Inches(0.78))
    line(af, "داده نشت‌دار", 21, True, NAVY, first=True, align=PP_ALIGN.CENTER, after=3)
    line(af, "۰٫۸۹۶ کنار گذاشته شد", 19, False, SLATE, align=PP_ALIGN.CENTER, after=0)
    c, cf = textbox(sl, Inches(4.66), Inches(5.72), Inches(3.85), Inches(0.78))
    line(cf, "محدودیت حافظه", 21, True, NAVY, first=True, align=PP_ALIGN.CENTER, after=3)
    line(cf, "توقف run6", 19, False, SLATE, align=PP_ALIGN.CENTER, after=0)
    b, bf = textbox(sl, MARGIN, Inches(5.72), Inches(3.85), Inches(0.78))
    line(bf, "مدل نهایی", 21, True, NAVY, first=True, align=PP_ALIGN.CENTER, after=3)
    line(bf, "۰٫۸۷۴ روی آزمون", 19, False, SLATE, align=PP_ALIGN.CENTER, after=0)
    animate(sl, [a, c, b])


# ---------------------------------------------------------------------- 11
@reg
def flaws():
    sl = slide("چالش‌های مجموعه‌داده")
    items = ["نشت داده: پراکندگی فریم‌های یک ویدیو میان سه بخش",
             "ویدیوی تکراری: یک مسیر با دو شناسه",
             "برچسب معیوب: کادرهای جابه‌جا، ناقص یا ناموجود",
             "ناهم‌ترازی ارزیابی: تفاوت توزیع اعتبارسنجی و آزمون"]
    tb, tf = textbox(sl, Inches(5.65), TOP + Inches(0.18), Inches(7.05), Inches(3.85))
    for i, txt in enumerate(items):
        line(tf, txt, 25, False, INK, first=(i == 0), dot="•", after=20)
    end, ef = textbox(sl, MARGIN, Inches(5.28), CW, Inches(0.78), MSO_ANCHOR.MIDDLE)
    line(ef, "خروجی مرحله: مجموعه‌داده آماده برای آموزش نهایی", 25, True,
         NAVY, first=True, align=PP_ALIGN.CENTER, after=0)
    animate(sl, [tb, end])


# ---------------------------------------------------------------------- 12
@reg
def leakage():
    sl = slide("نشت داده در تقسیم‌بندی اولیه")
    tb, tf = textbox(sl, Inches(6.45), TOP, Inches(6.25), Inches(0.58))
    line(tf, "نمونه واقعی از Record001_D، فریم‌های ۰ تا ۱۱", 24, True,
         NAVY, first=True, after=0)
    groups = [
        ["فریم ۰: آموزش", "فریم ۱: آموزش", "فریم ۲: آزمون", "فریم ۳: آموزش"],
        ["فریم ۴: آزمون", "فریم ۵: آزمون", "فریم ۶: آموزش", "فریم ۷: آموزش"],
        ["فریم ۸: آموزش", "فریم ۹: آموزش", "فریم ۱۰: آموزش", "فریم ۱۱: اعتبارسنجی"],
    ]
    x = Inches(8.65)
    marks = []
    for group in groups:
        box_text, frame = textbox(sl, x, Inches(2.12), Inches(3.55), Inches(2.70))
        for i, txt in enumerate(group):
            line(frame, txt, 22, False, INK, first=(i == 0), dot="•", after=12)
        marks.append(box_text)
        x -= Inches(4.05)
    bad, bf = textbox(sl, Inches(6.85), Inches(5.25), Inches(5.85), Inches(0.92))
    line(bf, "پیامد", 23, True, NAVY, first=True, after=3)
    line(bf, "ارزیابی خوش‌بینانه به‌علت شباهت فریم‌های مجاور", 20, False, SLATE, after=0)
    good, gf = textbox(sl, MARGIN, Inches(5.25), Inches(5.85), Inches(0.92))
    line(gf, "راهکار", 23, True, NAVY, first=True, after=3)
    line(gf, "تخصیص کامل هر ویدیو به یک بخش", 20, False, SLATE, after=0)
    animate(sl, marks + [bad, good])


# ---------------------------------------------------------------------- 13
@reg
def cleaning():
    sl = slide("مراحل آماده‌سازی داده نهایی")
    main = picture(sl, os.path.join(F2, "fig_2_1.png"), TOP - Inches(0.05),
                   height=Inches(3.95))
    items = [("کنترل کیفیت", "تکراری و برچسب معیوب"),
             ("تقسیم گروهی", "کل ویدیو در یک بخش"),
             ("اعتبارسنجی", "اندازه، رده و محیط")]
    x = Inches(8.85)
    shapes = []
    for head, body in items:
        a, af = textbox(sl, x, Inches(5.42), Inches(3.85), Inches(0.82))
        line(af, head, 21, True, NAVY, first=True, align=PP_ALIGN.CENTER, after=3)
        line(af, body, 19, False, SLATE, align=PP_ALIGN.CENTER, after=0)
        shapes.append(a)
        x -= Inches(4.19)
    animate(sl, [main] + shapes)


# ---------------------------------------------------------------------- 14
@reg
def subset():
    sl = slide("ترکیب زیرمجموعه نهایی")
    rows = [["بخش", "تعداد ویدیو", "تعداد تصویر", "تعداد نمونه"],
            ["آموزش", "۷۹", "۲۶٬۰۰۰", "۲۰۴٬۰۰۹"],
            ["اعتبارسنجی", "۳۹", "۵٬۵۷۱", "۳۶٬۸۱۰"],
            ["آزمون", "۴۲", "۵٬۵۷۱", "۳۴٬۹۹۴"]]
    strong_table(sl, rows, MARGIN, TOP, CW, Inches(2.48), 22, 21, hcols=(0,))
    tb, tf = textbox(sl, Inches(6.95), Inches(4.28), Inches(5.75), Inches(0.95))
    line(tf, "تقسیم ۷۰ / ۱۵ / ۱۵", 24, True, NAVY, first=True, after=3)
    line(tf, "بر اساس تعداد تصویر", 20, False, SLATE, after=0)
    tb2, tf2 = textbox(sl, MARGIN, Inches(4.28), Inches(5.75), Inches(0.95))
    line(tf2, "۱۶۰ ویدیو", 24, True, NAVY, first=True, after=3)
    line(tf2, "بدون اشتراک ویدیو یا تصویر میان بخش‌ها", 20, False, SLATE, after=0)
    foot, ff = textbox(sl, MARGIN, Inches(5.55), CW, Inches(0.72))
    line(ff, "۱۶۶ شناسه برچسب‌دار؛ پنج حذف کیفی و یک مورد بدون برچسب؛ ۱۶۰ ویدیوی نهایی",
         21, False, SLATE, first=True, align=PP_ALIGN.CENTER, after=0)


# --------------------------------------------------------------------- 14b
@reg
def validity():
    sl = slide("هم‌ترازی توزیع اندازه اشیا")
    picture(sl, os.path.join(F2, "fig_2_3.png"), TOP - Inches(0.05),
            height=Inches(5.35))


# --------------------------------------------------------------------- 14c
@reg
def balance():
    sl = slide("هم‌ترازی رده‌ها و شرایط محیطی")
    picture(sl, os.path.join(F2, "fig_2_2.png"), TOP - Inches(0.02), width=Inches(6.18),
            left=Inches(6.73))
    picture(sl, os.path.join(F2, "fig_2_4.png"), TOP - Inches(0.02), width=Inches(6.18),
            left=Inches(0.42))


# --------------------------------------------------------------------- 13b
@reg
def audit():
    sl = slide("بازبینی ویدیوهای مشکوک")
    picture(sl, os.path.join(FS, "label_noise.png"), TOP + Inches(0.05),
            height=Inches(3.75), left=MARGIN)
    tb, tf = textbox(sl, Inches(7.18), TOP + Inches(0.05), Inches(5.52), Inches(3.95))
    lines = [("غربال‌گری", "۱۳ ویدیو با mAP@0.5 کمتر از ۰٫۶۳"),
             ("خرابی تأییدشده", "Record426_D: ۰٫۰۰۷\nRecord046_D: ۰٫۴۵۱"),
             ("نتیجه بازبینی", "۲ حذف، ۱۱ حفظ")]
    for i, (head, body) in enumerate(lines):
        line(tf, head, 25, True, NAVY, first=(i == 0), after=4)
        line(tf, body, 21, False, SLATE, after=16)
    cap, cf = textbox(sl, MARGIN, Inches(5.30), Inches(6.35), Inches(0.72))
    line(cf, "نمونه Record426_D: کادر نادرست روی ساختمان", 20, True,
         CRIMSON, first=True, align=PP_ALIGN.CENTER, after=0)
    note, nf = textbox(sl, Inches(7.18), Inches(5.12), Inches(5.52), Inches(1.00))
    line(nf, "ملاک حذف", 23, True, NAVY, first=True, after=3)
    line(nf, "خرابی تأییدشده برچسب، نه امتیاز پایین مدل", 20, False, SLATE, after=0)
    animate(sl, [tb, note])
