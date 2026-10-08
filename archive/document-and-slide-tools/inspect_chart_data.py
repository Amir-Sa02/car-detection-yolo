from pathlib import Path
import json, numpy as np, pandas as pd
from docx import Document

ROOT=Path(r'C:\Users\amir2\Desktop\car-detection-yolo')
npz=np.load(ROOT/'figure and table data/eval_curves.npz')
print('NPZ keys:',npz.files)
for k in npz.files:
 a=npz[k]
 print(k,a.shape, str(a)[:200] if a.size<15 else '')
print('confusion:\n',npz['confusion'])
for name in ['run8/run8','run8/run8_ft','run7/run7']:
 folder=ROOT/'runs'/name
 for f in folder.glob('results*.csv'):
  d=pd.read_csv(f);d.columns=d.columns.str.strip()
  print(name,f.name,'rows',len(d))
  for key in ['metrics/mAP50(B)','metrics/mAP50-95(B)']:
   if key in d:
    i=d[key].idxmax();print('best',key,d.loc[i,['epoch',key]].to_dict())
  print('first,last',d.iloc[[0,-1]][['epoch','lr/pg0','metrics/mAP50(B)','metrics/mAP50-95(B)']].to_dict('records'))
doc=Document(r'C:\Users\amir2\Desktop\نسخه ارسالی\documents\پایان‌نامه.docx')
txt='\n'.join(p.text for p in doc.paragraphs)
Path('charts_thesis_current.txt').write_text(txt,encoding='utf-8')
text=txt[txt.find('۳\u200f-\u200f۱- روند آموزش'):]
print('THESIS chart captions:')
for p in doc.paragraphs:
 if p.text.startswith('شکل ') and not p.style.name.startswith('TOC'):print(p.text)
