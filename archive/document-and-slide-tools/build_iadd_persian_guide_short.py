from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(r"C:\Users\amir2\Desktop\نسخه ارسالی\مطالعه مقاله IADD\راهنمای_فارسی_مطالعه_مقاله_IADD.docx")
FA, EN = "B Zar", "Times New Roman"
NAVY, BLUE, GREEN, ORANGE = "123A63", "2C73A8", "2F7D4A", "B75A16"
INK, WHITE, GRAY = "18222E", "FFFFFF", "F1F4F6"
PALE_BLUE, PALE_GREEN, PALE_ORANGE = "EAF3F9", "EAF5EC", "FFF1E8"
MARKS = re.compile(r"[\u064B-\u065F\u0670\u06D6-\u06ED]")
LATIN = re.compile(r"([A-Za-z][A-Za-z0-9_.@:+/%-]*(?:[ \t]+[A-Za-z][A-Za-z0-9_.@:+/%-]*)*)")


def clean(s):
    return MARKS.sub("", s)


def rtl(p):
    ppr = p._p.get_or_add_pPr()
    b = ppr.find(qn("w:bidi"))
    if b is None:
        b = OxmlElement("w:bidi")
        ppr.append(b)
    b.set(qn("w:val"), "1")


def run_props(run, latin=False, size=13.5, bold=False, color=INK):
    run.bold = bold
    run.font.name = EN if latin else FA
    run.font.size = Pt(size - 1 if latin else size)
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    name = EN if latin else FA
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)
    r = rpr.find(qn("w:rtl"))
    if r is None:
        r = OxmlElement("w:rtl")
        rpr.append(r)
    r.set(qn("w:val"), "0" if latin else "1")


def mixed(p, text, size=13.5, bold=False, color=INK):
    for part in LATIN.split(clean(text)):
        if not part:
            continue
        rr = p.add_run(part)
        run_props(rr, bool(LATIN.fullmatch(part)), size, bold, color)


def para(doc, text, size=13.5, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         before=0, after=3, line=1.08, keep=False):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    mixed(p, text, size, bold, color)
    return p


def heading(doc, text, level=1):
    p = para(doc, text, 18 if level == 1 else 15, True, NAVY if level == 1 else BLUE,
             WD_ALIGN_PARAGRAPH.RIGHT, 4 if level == 1 else 2, 4, 1.0, True)
    if level == 1:
        ppr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), BLUE)
        pbdr.append(bottom)
        ppr.append(pbdr)
    return p


def bullet(doc, text, size=13.2, color=INK):
    p = para(doc, "• " + text, size, color=color, align=WD_ALIGN_PARAGRAPH.RIGHT, after=2, line=1.05)
    p.paragraph_format.right_indent = Cm(0.25)
    p.paragraph_format.first_line_indent = Cm(-0.2)
    return p


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    s = tcpr.find(qn("w:shd"))
    if s is None:
        s = OxmlElement("w:shd")
        tcpr.append(s)
    s.set(qn("w:fill"), fill)


def margins(cell, value=90):
    tcpr = cell._tc.get_or_add_tcPr()
    tcm = tcpr.first_child_found_in("w:tcMar")
    if tcm is None:
        tcm = OxmlElement("w:tcMar")
        tcpr.append(tcm)
    for name in ("top", "start", "bottom", "end"):
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:w"), str(value))
        e.set(qn("w:type"), "dxa")
        tcm.append(e)


def table(doc, headers, rows, widths, size=11.6):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bidi = OxmlElement("w:bidiVisual")
    bidi.set(qn("w:val"), "1")
    t._tbl.tblPr.append(bidi)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Cm(widths[i])
        shade(c, NAVY); margins(c)
        p = c.paragraphs[0]; rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mixed(p, h, size, True, WHITE)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]; c.width = Cm(widths[i]); margins(c); shade(c, WHITE if ri % 2 == 0 else GRAY)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]; rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            mixed(p, str(val), size, False, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def box(doc, title, text, fill=PALE_BLUE, accent=BLUE, size=12.8):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0); shade(c, fill); margins(c, 130)
    p = c.paragraphs[0]; rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mixed(p, title + "\n", size + 0.4, True, accent)
    mixed(p, text, size, False, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def new_page(doc):
    doc.add_page_break()


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.55), Cm(1.45)
sec.right_margin, sec.left_margin = Cm(1.75), Cm(1.75)
sec.header_distance, sec.footer_distance = Cm(0.7), Cm(0.65)
st = doc.styles["Normal"]
st.font.name, st.font.size = FA, Pt(13.5)
st._element.rPr.rFonts.set(qn("w:cs"), FA)

# Footer.
fp = sec.footer.paragraphs[0]; rtl(fp); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
mixed(fp, "صفحه ", 10.5, False, "66707A")
r = OxmlElement("w:r"); fb = OxmlElement("w:fldChar"); fb.set(qn("w:fldCharType"), "begin")
it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
fs = OxmlElement("w:fldChar"); fs.set(qn("w:fldCharType"), "separate")
tx = OxmlElement("w:t"); tx.text = "1"
fe = OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"), "end")
r.extend([fb, it, fs, tx, fe]); fp._p.append(r)

# Page 1: compact overview.
para(doc, "راهنمای فشرده مطالعه مقاله اصلی IADD", 23, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 2, 2, 1.0)
para(doc, "مطالب ضروری برای دفاع در ۲ تا ۳ ساعت مطالعه", 14, False, GREEN, WD_ALIGN_PARAGRAPH.CENTER, 0, 7, 1.0)
box(doc, "شناسنامه مقاله", "Khosravian و همکاران، دانشکده مهندسی خودرو دانشگاه علم و صنعت ایران، نشریه IET Image Processing، جلد ۱۷، سال ۲۰۲۳، DOI: 10.1049/ipr2.12710", PALE_BLUE, NAVY)
heading(doc, "خلاصه ای که باید بلد باشید", 1)
para(doc, "مقاله IADD برای کاهش افت عملکرد آشکارسازها در شهرها و شرایط جدید، یک مجموعه داده چنددامنه از رانندگی واقعی در ایران ارائه می کند. داده ها از هشت شهر، مسیرهای شهری و برون شهری، روز، شب و باران و چند تفکیک پذیری جمع آوری شده اند. شش رده مقاله عبارت اند از شخص، خودروی سبک، موتورسیکلت، اتوبوس، خودروی باری و چراغ راهنمایی.")
para(doc, "مقاله ۹۷٬۵۲۸ داده برچسب دار از ۱۷۱ ویدیو گزارش می کند. حدود ۳۰٬۰۰۰ تصویر ابتدا دستی برچسب خوردند؛ سپس YoloR بقیه تصاویر را برچسب زد و کارشناسان خروجی ها را بازبینی کردند. نویسندگان با Faster R-CNN، EfficientDetB4، YOLOv4 و RetinaNet نشان دادند که تنوع IADD می تواند تعمیم به محیط های جدید را بهتر کند.")
box(doc, "یک جمله دفاعی", "مقاله منبع معرفی IADD و روش ساخت اولیه آن است؛ اما کشف مشکلات نسخه منتشرشده و ساخت تقسیم ویدیویی نهایی، کار خود پروژه ماست.", PALE_GREEN, GREEN)

# Page 2: data and labeling.
new_page(doc)
heading(doc, "۱. داده و برچسب گذاری", 1)
table(doc, ["موضوع", "اطلاعات مقاله"], [
    ["شهرها", "تهران، البرز، اصفهان، بوشهر، ملایر، اراک، ساوه و رشت"],
    ["محیط", "بزرگراه، خیابان، کوچه، میدان، محیط روستایی، ترافیک روان و متراکم"],
    ["شرایط", "روز، شب و باران"],
    ["تفکیک پذیری", "۶۴۰×۴۸۰ تا ۳۸۴۰×۲۱۶۰ پیکسل"],
    ["استخراج فریم", "نیم فریم بر ثانیه از ویدیوهای ۳۰ یا ۶۰ فریم بر ثانیه"],
    ["رده ها", "person، car، motorcycle، bus، truck و traffic lights"],
], [5.0, 11.8])
heading(doc, "فرایند برچسب گذاری", 2)
for s in [
    "برچسب گذاری دستی حدود ۳۰٬۰۰۰ تصویر",
    "آموزش مدل YoloR با تصاویر دستی",
    "تولید خودکار کادرهای تصاویر باقی مانده",
    "بازبینی و اصلاح خروجی ها توسط کارشناسان",
    "ذخیره هر برچسب در قالب YOLO: شناسه رده، مرکز کادر، عرض و ارتفاع نرمال شده",
]: bullet(doc, s)
box(doc, "نکته درباره نویز", "مقاله می گوید خروجی ها بازبینی شده اند، اما ادعا نمی کند همه برچسب ها بدون خطا هستند. کادرهای گمشده، ناقص یا جابه جا را خود پروژه در نسخه منتشرشده مشاهده کرده است.", PALE_ORANGE, ORANGE)
heading(doc, "چرا IADD برای پروژه مناسب بود؟", 2)
para(doc, "این مجموعه داده صحنه های واقعی رانندگی ایران را پوشش می دهد؛ بنابراین از نظر تابلوها، خیابان ها، وسایل نقلیه و شرایط محیطی به دامنه هدف پروژه نزدیک تر از مجموعه ای مانند BDD100K بود.")

# Page 3: official vs project.
new_page(doc)
heading(doc, "۲. تفاوت آمار مقاله با نسخه مورد استفاده پروژه", 1)
table(doc, ["موضوع", "مقاله", "یافته پروژه"], [
    ["کل داده", "۹۷٬۵۲۸ داده برچسب دار", "۸۳٬۰۴۲ فایل برچسب در train و val"],
    ["ویدیو", "۱۷۱ ویدیو", "ویدیوهای قابل استفاده پس از حذف موارد مسئله دار کمتر شدند"],
    ["test رسمی", "۱۴٫۸٪ از تقسیم رسمی", "۱۴٬۴۸۶ تصویر، بدون فایل برچسب عمومی"],
    ["تقسیم", "۷۸٪ آموزش، ۷٫۲٪ اعتبارسنجی، ۱۴٫۸٪ آزمون", "تقسیم جدید در سطح کل ویدیو"],
    ["داده بدون برچسب", "حدود ۷٫۵ هزار داده جداگانه در جدول مقایسه", "با ۱۴٬۴۸۶ تصویر test یکی فرض نشود"],
], [4.1, 5.7, 7.0], 11.2)
box(doc, "رابطه عددها", "۸۳٬۰۴۲ تصویر دارای فایل برچسب + ۱۴٬۴۸۶ تصویر test بدون فایل برچسب عمومی = ۹۷٬۵۲۸ تصویر گزارش شده در مقاله.", PALE_GREEN, GREEN)
heading(doc, "مطالبی که مقاله نگفته است", 2)
for s in [
    "نبود فایل برچسب عمومی test: مشاهده مستقیم در پوشه های دانلودشده",
    "پراکندگی فریم های یک ویدیو میان بخش ها: نتیجه تحلیل شناسه ویدیوها",
    "سه ویدیوی تکراری و دو ویدیوی دارای برچسب معیوب: یافته پروژه",
    "معنای کدهای D، N، R و A: مقاله آن ها را تعریف نمی کند",
    "تقسیم ویدیویی، هم ترازی دشواری و نمونه گیری زیرمجموعه نهایی: طراحی پروژه",
]: bullet(doc, s, 12.8)
box(doc, "پاسخ دفاعی", "عدد ۹۷٬۵۲۸ از مقاله آمده است؛ اما نبود برچسب test و مشکلات تقسیم رسمی را با بررسی فایل های منتشرشده فهمیدیم. به همین دلیل داده دارای برچسب را در سطح کل ویدیو دوباره تقسیم کردیم.", PALE_BLUE, BLUE)

# Page 4: experiments and results.
new_page(doc)
heading(doc, "۳. آزمایش ها و نتایج مقاله", 1)
heading(doc, "سه سناریوی اصلی", 2)
table(doc, ["سناریو", "آموزش", "ارزیابی", "هدف"], [
    ["۱", "KITTI", "IADD", "بررسی انتقال به دامنه ایران"],
    ["۲", "IADD", "IADD", "ارزیابی مدل ها روی IADD"],
    ["۳", "KITTI + IADD", "IADD", "اثر افزودن تنوع دامنه"],
], [2.0, 3.8, 3.8, 7.2])
para(doc, "همه مدل ها با وزن های از پیش آموزش دیده روی COCO آغاز شدند. معیار اصلی مقاله AP و mAP است؛ AP مساحت زیر منحنی دقت برحسب فراخوانی و mAP میانگین AP رده هاست.")
heading(doc, "نتیجه YOLOv4 در سناریوی دوم", 2)
table(doc, ["شخص", "خودروی سبک", "موتورسیکلت", "اتوبوس", "خودروی باری", "چراغ", "mAP"], [["۸۶٫۶", "۹۶٫۱", "۸۳٫۴", "۸۱٫۵", "۹۰٫۱", "۸۱٫۹", "۸۶٫۶"]], [2.35]*7, 10.5)
box(doc, "چرا با ۰٫۸۷۴ پروژه مقایسه مستقیم نکردیم؟", "داده آزمون و تقسیم یکسان نیست، مدل و کد و تنظیمات متفاوت اند و مقاله آستانه IoU عدد ۸۶٫۶ را در همان بخش به روشنی مشخص نکرده است. بنابراین این دو عدد، آزمایش کنترل شده روی شرایط برابر نیستند.", PALE_ORANGE, ORANGE)
heading(doc, "آزمایش CARLA", 2)
para(doc, "مقاله مدل های آموزش دیده با IADD و KITTI را روی تصاویر شبیه سازی شده CARLA آزمایش کرد. عملکرد بهتر مدل های آموزش دیده با IADD به عنوان شاهدی برای اثر تنوع دامنه مطرح شد. CARLA در پروژه شما استفاده نشده و نتیجه آن نباید به YOLOv11 پروژه نسبت داده شود.")

# Page 5: mapping and critical reading.
new_page(doc)
heading(doc, "۴. کدام بخش ها در پایان نامه و ارائه استفاده شده اند؟", 1)
table(doc, ["مطلب", "صفحه PDF مقاله", "کاربرد"], [
    ["هدف IADD و تعمیم", "۱ و ۲", "دلیل انتخاب مجموعه داده"],
    ["۹۷٬۵۲۸ داده، ۱۷۱ ویدیو، شش رده", "۱، ۲ و ۶", "معرفی IADD"],
    ["شهرها، مسیرها و شرایط", "۲ و ۵", "توضیح تنوع داده"],
    ["نیم فریم بر ثانیه و تفکیک پذیری", "۵", "مشخصات تصاویر"],
    ["۳۰ هزار تصویر و YoloR", "۵ و ۶", "فرایند برچسب گذاری"],
    ["پیش آموزش COCO", "۸", "یادگیری انتقالی"],
    ["تعریف AP و mAP", "۸", "معیارهای ارزیابی"],
    ["نتیجه YOLOv4", "۹", "آمادگی برای سوال مقایسه"],
    ["CARLA و جمع بندی", "۱۲ و ۱۳", "شواهد تعمیم مقاله"],
], [5.2, 3.1, 8.5], 11.2)
heading(doc, "پنج نکته ای که نباید اشتباه بگویید", 2)
for s in [
    "مقاله نگفته است بخش test منتشرشده برچسب ندارد؛ این مشاهده پروژه است.",
    "مقاله نشت فریم میان بخش ها را گزارش نکرده است؛ این مشکل را پروژه کشف کرد.",
    "مقاله وجود سه ویدیوی تکراری و دو ویدیوی معیوب را گزارش نکرده است.",
    "بازبینی کارشناسی به معنی تضمین صفر بودن نویز برچسب نیست.",
    "mAP برابر ۸۶٫۶ مقاله، خط مبنای قابل مقایسه مستقیم با ۰٫۸۷۴ پروژه نیست.",
]: bullet(doc, s, 12.6)
box(doc, "نتیجه اصلی مقاله", "تنوع بیشتر شهر، مسیر، نور، آب وهوا و ابزار تصویربرداری می تواند ویژگی های عمومی تری به مدل بیاموزد و تعمیم آن را در محیط جدید بهتر کند.", PALE_GREEN, GREEN)

# Page 6: oral prep.
new_page(doc)
heading(doc, "۵. پاسخ های کوتاه برای دفاع", 1)
qas = [
    ("چرا IADD؟", "داده واقعی ایران با تنوع شهر، مسیر، نور و آب وهوا و نزدیک به دامنه هدف پروژه است."),
    ("برچسب گذاری چگونه بود؟", "۳۰ هزار تصویر دستی، آموزش YoloR، تولید خودکار بقیه کادرها و بازبینی کارشناسان."),
    ("چرا همه تصاویر را استفاده نکردید؟", "برای ۱۴٬۴۸۶ تصویر test فایل برچسب عمومی وجود نداشت؛ بدون حقیقت مبنا قابل استفاده نبودند."),
    ("چرا تقسیم جدید ساختید؟", "برای جلوگیری از حضور فریم های یک ویدیو در چند بخش و به دست آوردن ارزیابی مستقل."),
    ("چرا مقایسه مستقیم با YOLOv4 نکردید؟", "داده، تقسیم، مدل، کد، تنظیمات و تعریف دقیق معیار برابر نبودند."),
    ("منظور از تعمیم چیست؟", "حفظ عملکرد هنگام تغییر شهر، نور، آب وهوا، دوربین یا محیط نسبت به داده آموزشی."),
]
for q, a in qas:
    para(doc, q, 13.2, True, NAVY, WD_ALIGN_PARAGRAPH.RIGHT, 1, 0, 1.0, True)
    para(doc, a, 12.5, after=2, line=1.0)
heading(doc, "متن یک دقیقه ای", 2)
box(doc, "متن پیشنهادی", "مجموعه داده اصلی پروژه IADD است که پژوهشگران دانشگاه علم و صنعت ایران برای بهبود تعمیم آشکارسازها معرفی کرده اند. مقاله ۹۷٬۵۲۸ داده برچسب دار از ۱۷۱ ویدیو و شش رده را گزارش می کند. حدود ۳۰ هزار تصویر دستی برچسب خوردند و YoloR برای تولید برچسب بقیه تصاویر به کار رفت. ما در نسخه منتشرشده متوجه شدیم بخش test فایل برچسب عمومی ندارد و فریم های یک ویدیو میان بخش ها پراکنده اند. بنابراین فقط داده دارای برچسب را استفاده کردیم، موارد تکراری و معیوب را حذف کردیم و تقسیم جدیدی در سطح کل ویدیو ساختیم. نتیجه YOLOv4 مقاله با نتیجه پروژه مقایسه مستقیم نشد، چون شرایط آزمایش یکسان نبود.", PALE_BLUE, NAVY, 12.1)
para(doc, "اگر این صفحه و جدول صفحه ۳ را بدون نگاه کردن توضیح بدهید، برای سوال های اصلی مربوط به مقاله آماده هستید.", 12, True, GREEN, WD_ALIGN_PARAGRAPH.CENTER, 4, 0, 1.0)

doc.core_properties.title = "راهنمای فشرده مطالعه مقاله IADD"
doc.core_properties.author = "محمد امیر صادق زاده"
doc.save(OUT)
print(OUT.name.encode("unicode_escape").decode())
