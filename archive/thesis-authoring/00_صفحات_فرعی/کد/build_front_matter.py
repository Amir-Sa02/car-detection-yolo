# -*- coding: utf-8 -*-
"""روی جلد و صفحه‌های فرعی، بر پایه آیین‌نامه ۱۳۸۷ (بندهای ۱-۱ و ۲-۱، ص ۱۲ تا ۱۵).

هر اندازه‌ای که اینجا به کار رفته، کنار همان خط از آیین‌نامه نقل شده است.
فاصله‌ها در آیین‌نامه از لبه کاغذ داده شده‌اند، پس صفحه عنوان حاشیه‌های خودش را
دارد و هر بلوک با یک مکان‌نمای مطلق سر جای دقیقش می‌نشیند.

ترتیب صفحه‌های فرعی (ص ۱۴):
  ۱ بسم الله   ۲ صفحه عنوان   ۳ سپاسگزاری (اختیاری)   ۴ تقدیم
  ۵ چکیده و واژه‌های کلیدی   ۶ فهرست مطالب   ۷ فهرست شکل‌ها   ۸ فهرست جدول‌ها
  ۹ لیست علائم و اختصارات
«صفحه‌های فرعی فقط بر اساس حروف الفبای فارسی شماره‌گذاری می‌شود» و Word قالب
شماره‌گذاری فارسی ندارد، پس هر صفحه یک سکشن جداست و حرفش در فوتر نوشته شده است.

سه فهرست، فیلد واقعی Word هستند و پس از ادغام با فصل‌ها با Ctrl+A و F9 پر می‌شوند.

دو تله‌ای که در نسخه‌های پیشین این فایل خرابی به بار آورد:
  * در پاراگراف راست‌به‌چپ، Word مقدارهای قدیمی left/right را آینه می‌کند؛
    پس هر تراز و هر تورفتگی اینجا با start/end نوشته شده است.
  * فاصله سطر EXACTLY هر چیزی بلندتر از خودش را می‌بُرد - آرم را نصف می‌کرد و
    حروف فارسی را توی هم می‌ریخت. برای متن و تصویر از AT_LEAST استفاده می‌شود و
    EXACTLY فقط برای فاصله‌گذارهای خالی می‌ماند.
"""
import os

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.normpath(os.path.join(HERE, "..", "متن"))
PICS = os.path.normpath(os.path.join(HERE, "..", "..", "pics"))
os.makedirs(OUT, exist_ok=True)

from docx import Document                                              # noqa: E402
from docx.shared import Pt, Cm, RGBColor                               # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING         # noqa: E402
from docx.enum.section import WD_SECTION                               # noqa: E402
from docx.oxml.ns import qn                                            # noqa: E402
from docx.oxml import OxmlElement                                      # noqa: E402
from PIL import Image                                                  # noqa: E402

# ============================================================ CONFIG
CONFIG = {
    "مقطع_رشته":    "پایان‌نامه دوره کارشناسی مهندسی کامپیوتر",
    # «دقت شود که عنوان دقیقاً مطابق عنوانی باشد که در هنگام انتخاب پروژه تأیید شده»
    "عنوان":        "[عنوان مصوب پروژه را اینجا بنویسید]",
    "عنوان_موقت":   True,                      # قرمز، تا وقتی عنوان مصوب معلوم شود
    "استاد_راهنما": "دکتر رضا شمسایی",
    "استاد_داور":   "[نام استاد داور]",
    "داور_موقت":    True,
    "دانشجو":       "محمد امیر صادق‌زاده",
    "تاریخ":        "شهریور ۱۴۰۵",
}

FONT, TITLE_FONT, LAT = "B Zar", "B Lotus", "Times New Roman"
RED = RGBColor(0xC0, 0, 0)
LETTERS = ["آ", "ب", "پ", "ت", "ث", "ج", "چ", "ح", "خ"]
CM = 28.3465                                   # points per centimetre
TWIP = 567                                     # twips per centimetre

doc = Document()
_st = doc.styles["Normal"]
_st.font.name = FONT; _st.font.size = Pt(14)
_st.paragraph_format.line_spacing = 1.5; _st.paragraph_format.space_after = Pt(0)


def _rs(run, size, bold=False, font=None):
    """python-docx only writes w:rFonts/@ascii; without w:cs Word silently falls
    back to another face for the Persian glyphs."""
    name = font or FONT
    rPr = run._r.get_or_add_rPr()
    f = rPr.find(qn("w:rFonts"))
    if f is None:
        f = OxmlElement("w:rFonts"); rPr.append(f)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        f.set(qn(a), name)
    szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size * 2))); rPr.append(szcs)
    if bold:
        rPr.append(OxmlElement("w:bCs"))


def _bidi(p):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.insert(0, OxmlElement("w:bidi"))


def _jc(p, val):
    """Word mirrors the legacy left/right values inside an RTL paragraph, so a
    «right-aligned» Persian line comes out flush LEFT. start/end never mirror:
    start is the right edge in an RTL paragraph."""
    pPr = p._p.get_or_add_pPr()
    for e in pPr.findall(qn("w:jc")):
        pPr.remove(e)
    e = OxmlElement("w:jc"); e.set(qn("w:val"), val); pPr.append(e)


def _ind(p, start_cm=0.0, end_cm=0.0):
    """same trap as _jc: w:left/w:right mirror in RTL, w:start/w:end do not"""
    pPr = p._p.get_or_add_pPr()
    for e in pPr.findall(qn("w:ind")):
        pPr.remove(e)
    if not (start_cm or end_cm):
        return
    e = OxmlElement("w:ind")
    e.set(qn("w:start"), str(int(start_cm * TWIP)))
    e.set(qn("w:end"), str(int(end_cm * TWIP)))
    pPr.append(e)


def para(text="", size=14, bold=False, jc="center", font=None, sa=0,
         start=0.0, end=0.0, height=None, color=None):
    """`height` is an AT_LEAST line height in cm - predictable for the page
    maths, but it grows instead of clipping when a glyph needs more room."""
    p = doc.add_paragraph()
    _bidi(p)
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(sa)
    if height:
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        pf.line_spacing = Pt(height * CM)
    else:
        pf.line_spacing = 1.5
    if text:
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
        r.font.name = font or FONT
        if color is not None:
            r.font.color.rgb = color
        _rs(r, size, bold, font)
    _ind(p, start, end)
    _jc(p, jc)
    return p


def gap(cm):
    """empty paragraph of an exact height - the vertical ruler of this file"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(max(0.12, cm) * CM)
    return p


def picture(fname, width_cm=None, height_cm=None):
    """returns the height it occupied, so the caller's cursor stays honest"""
    path = os.path.join(PICS, fname)
    if not os.path.exists(path):
        raise SystemExit("missing picture: " + fname)
    w_px, h_px = Image.open(path).size
    if height_cm is None:
        height_cm = width_cm * h_px / w_px
    if width_cm is None:
        width_cm = height_cm * w_px / h_px
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST      # never clip the image
    pf.line_spacing = Pt(height_cm * CM)
    p.add_run().add_picture(path, width=Cm(width_cm), height=Cm(height_cm))
    return height_cm * 1.04


def _page_setup(s, top=3.0, bottom=2.5):
    """A4، حاشیه راست ۳، چپ ۲، بالا ۳، پایین ۲٫۵، فوتر ۱٫۵ (ص ۱۵)"""
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
    s.right_margin, s.left_margin = Cm(3), Cm(2)
    s.footer_distance = Cm(1.5)
    if s._sectPr.find(qn("w:bidi")) is None:
        s._sectPr.append(OxmlElement("w:bidi"))


FA_LETTERS = ["آ", "ب", "پ", "ت", "ث", "ج", "چ", "ح", "خ",
              "د", "ذ", "ر", "ز", "ژ", "س", "ش"]


def _fld(p, kind, val):
    r = p.add_run(); r.font.size = Pt(12); _rs(r, 12)
    e = OxmlElement(kind)
    if kind == "w:fldChar":
        e.set(qn("w:fldCharType"), val)
    else:
        e.set(qn("xml:space"), "preserve"); e.text = val
    r._r.append(e)


def _letter_field(p, i=0):
    """{ IF { PAGE } = 1 "آ" { IF { PAGE } = 2 "ب" ... } }

    یک حرف برای هر صفحه، نه برای هر بخش. اگر فهرست مطالب دو یا سه صفحه شود،
    صفحه‌های بعدی حرف بعدی خود را می‌گیرند و شماره‌گذاری نمی‌لنگد. Word هیچ
    قالب شماره‌گذاری با حروف فارسی ندارد (arabicAlpha حرف‌های پ و چ را ندارد)،
    پس نگاشت شماره صفحه به حرف باید با فیلد ساخته شود."""
    if i >= len(FA_LETTERS):
        _fld(p, "w:instrText", ' "" ')
        return
    _fld(p, "w:fldChar", "begin")
    _fld(p, "w:instrText", " IF ")
    _fld(p, "w:fldChar", "begin")
    _fld(p, "w:instrText", " PAGE ")
    _fld(p, "w:fldChar", "end")
    _fld(p, "w:instrText", ' = %d "%s" ' % (i + 1, FA_LETTERS[i]))
    _letter_field(p, i + 1)
    _fld(p, "w:fldChar", "end")


def _footer_letter(s, letter):
    s.footer.is_linked_to_previous = False
    p = s.footer.paragraphs[0]
    p.text = ""
    _bidi(p); _jc(p, "center")
    if letter:
        _letter_field(p)


def new_section(letter, top=3.0, bottom=2.5):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    _page_setup(s, top, bottom)
    _footer_letter(s, letter)
    return s


def _field(p, instr):
    """a live Word field, so the three فهرست pages fill themselves in"""
    for tag, attr, txt in (("w:fldChar", "w:fldCharType", "begin"),
                           ("w:instrText", None, instr),
                           ("w:fldChar", "w:fldCharType", "separate"),
                           (None, None, "برای به‌روزرسانی این فهرست: Ctrl+A سپس F9"),
                           ("w:fldChar", "w:fldCharType", "end")):
        r = p.add_run()
        if tag is None:
            r.text = txt; r.font.size = Pt(12); _rs(r, 12)
            continue
        e = OxmlElement(tag)
        if attr:
            e.set(qn(attr), txt)
        else:
            e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)


def title_block(with_judge):
    """آرم ۲ · مقطع و رشته ۷٫۵ · «موضوع:» ۱۰ · عنوان ۱۱٫۵      (از بالای کاغذ)
       استاد راهنما ۱۲ · نام دانشجو ۷ · تاریخ ۳٫۵               (از پایین کاغذ)
    «صفحه عنوان همان مندرجات روی جلد است، تنها با افزودن نام استاد داور» (ص ۱۴)
    فونت‌ها طبق ص ۱۲: مقطع و رشته ۱۶ زر سیاه، «موضوع:» ۱۸ زر سیاه، عنوان ۲۰ زر
    سیاه، استاد راهنما و نام دانشجو ۱۸ زر سیاه، تاریخ ۱۶ زر (بدون سیاه).
    """
    cur = [2.0]                                 # حاشیه بالای این صفحه

    def to(edge_cm):
        d = edge_cm - cur[0]
        if d > 0.05:
            gap(d); cur[0] = edge_cm

    def line(text, size, bold=True, color=None, start=0.0, end=0.0, lines=1):
        h = size * 1.7 / CM                     # جای راحت برای کشیدگی حروف فارسی
        para(text, size, bold, height=h, color=color, start=start, end=end)
        cur[0] += lines * h

    cur[0] += picture("لوگو_جدید_دانشگاه.png", height_cm=3.4)
    to(7.5)
    line(CONFIG["مقطع_رشته"], 16)
    to(10.0)
    line("موضوع:", 18)
    to(11.5)
    per_line = max(1, int(13.0 * CM / (20 * 0.46)))
    n_lines = max(1, -(-len(CONFIG["عنوان"]) // per_line))
    line(CONFIG["عنوان"], 20, start=1.5, end=1.5, lines=n_lines,
         color=RED if CONFIG["عنوان_موقت"] else None)
    to(29.7 - 12.0)                             # استاد راهنما، ۱۲cm از پایین
    line("استاد راهنما:", 18)
    line(CONFIG["استاد_راهنما"], 18)
    if with_judge:
        gap(0.5); cur[0] += 0.5
        line("استاد داور:", 18)
        line(CONFIG["استاد_داور"], 18, color=RED if CONFIG["داور_موقت"] else None)
    to(29.7 - 7.0)                              # نام دانشجو، ۷cm از پایین
    line("نام دانشجو:", 18)
    line(CONFIG["دانشجو"], 18)
    to(29.7 - 3.5)                              # تاریخ، ۳٫۵cm از پایین
    line(CONFIG["تاریخ"], 16, bold=False)



def build_front(target, first_section_ready=False):
    """بسازد صفحه‌های فرعی را داخل سندی که به آن داده می‌شود.

    `first_section_ready` یعنی سکشن جاری همین حالا برای صفحه بسم الله آماده است
    (حالت ادغام: قبلش چیزی در سند نیست)."""
    global doc
    doc = target
    # ==================================================== ۱. بسم الله (صفحه آ)
    _page_setup(doc.sections[0])
    _footer_letter(doc.sections[0], LETTERS[0])
    gap(8.0)
    picture("بسم_الله_برش‌خورده.png", width_cm=13.0)


    # ==================================== ۲. صفحه عنوان (صفحه ب) و روی جلد
    new_section(LETTERS[1], top=2.0, bottom=1.5)
    title_block(with_judge=True)

    # ==================================================== ۳. سپاسگزاری (صفحه پ)
    # «تیتر سپاسگزاری ۹٫۵cm پایین‌تر از بالای صفحه و از دو طرف کاملاً وسط صفحه و با
    #  فونت ۱۶ نوشته شود و مطالب آن با فونت ۱۴ با فاصله ۴cm از سمت راست و ۳cm از
    #  سمت چپ نوشته شود» (ص ۱۴). حاشیه راست ۳ و چپ ۲ است، پس start=۱ و end=۱.
    new_section(LETTERS[2])
    gap(6.5)                                        # ۳ + ۶٫۵ = ۹٫۵
    para("سپاسگزاری", 16, True)
    gap(1.0)
    para("از استاد راهنمای گرامی‌ام، جناب آقای دکتر رضا شمسایی، که در تمام مراحل این پژوهش، "
         "از طرح مسئله تا نگارش نهایی، با راهنمایی‌های دقیق و شکیبایی بسیار مرا همراهی کردند، "
         "صمیمانه سپاسگزارم. همچنین از استادان گروه مهندسی کامپیوتر که در سال‌های تحصیل، دانش و "
         "تجربه خود را بی‌دریغ در اختیارم گذاشتند، و از خانواده‌ام که پشتیبان همیشگی من بوده‌اند، "
         "قدردانی می‌کنم.", 14, jc="both", start=1, end=1)

    # ==================================================== ۴. تقدیم (صفحه ت)
    # «تیتر «تقدیم به» با فونت ۱۲ نوشته و فاصله آن از بالا ۱۰cm و از سمت راست ۱۳cm
    #  باشد. مطالب این قسمت با فونت ۱۲ و با طول سطر ۵cm نوشته شود» (ص ۱۴).
    # ۱۳cm از لبه راست کاغذ یعنی x=۸؛ حاشیه راست ۳cm است، پس start=۱۰ و برای رسیدن
    # به طول سطر ۵cm مقدار end=۱ می‌شود.
    new_section(LETTERS[3])
    gap(7.0)                                        # ۳ + ۷ = ۱۰
    para("تقدیم به", 12, True, jc="start", start=10, end=1)
    gap(1.0)
    para("جویندگان دانش", 12, jc="start", start=10, end=1)
    para("که آموختن را پایانی نمی‌دانند", 12, jc="start", start=10, end=1)

    # ==================================================== ۵. چکیده (صفحه ث)
    # «عنوان «چکیده» با فونت ۱۴ لوتوس سیاه نسبت به بالای صفحه ۴cm است، کاملاً در وسط
    #  صفحه نوشته می‌شود و در سطر بعد و در وسط صفحه عنوان پروژه با همان اندازه.
    #  مطالب چکیده با فونت ۱۲ زر با فاصله ۴cm از سمت راست و ۳cm از سمت چپ و حداکثر
    #  در نصف صفحه نوشته شود» (ص ۱۴)
    new_section(LETTERS[4])
    gap(1.0)                                        # ۳ + ۱ = ۴
    para("چکیده", 14, True, font=TITLE_FONT)
    para(CONFIG["عنوان"], 14, True, font=TITLE_FONT, start=1, end=1,
         color=RED if CONFIG["عنوان_موقت"] else None)
    gap(1.0)
    para("در این پژوهش، یک آشکارساز اشیای ترافیکی بر پایه معماری YOLOv11 روی مجموعه‌داده رانندگی "
         "خودکار ایران آموزش داده شد و ارزیابی گردید. از آنجا که تصاویر این مجموعه از فریم‌های "
         "پیاپی ویدیو استخراج شده‌اند، پیش از آموزش چهار اشکال شناسایی و برطرف شد: امکان نشت داده "
         "میان بخش‌های آموزش و ارزیابی، وجود ویدیوهای تکراری با شناسه متفاوت، ویدیوهایی با "
         "برچسب معیوب، و ناهم‌ترازی دشواری دو بخش ارزیابی. تقسیم داده در سطح کل ویدیو و با هم‌ترازی توزیع اندازه اشیا، رده‌ها و شرایط "
         "محیطی انجام شد. سپس دو راهبرد بهینه‌سازی در شرایط یکسان مقایسه شدند و پیکربندی برتر "
         "برگزیده شد. مدل نهایی روی بخش آزمون به میانگین دقت متوسط ۰٫۸۷۴ در آستانه هم‌پوشانی ۰٫۵ "
         "دست یافت. تحلیل به تفکیک رده نشان داد که دشواری هر رده را باید بر پایه دو عامل تشابه "
         "بصری و اندازه شیء تفسیر کرد، نه صرفاً بر پایه شمار نمونه‌ها.",
         12, jc="both", start=1, end=1)
    gap(1.0)                                        # «حداقل با یک خط فاصله»
    # «عنوان واژه‌های کلیدی ... در سمت راست با فونت ۱۴ لوتوس سیاه»
    para("واژه‌های کلیدی", 14, True, font=TITLE_FONT, jc="start", start=1)
    para("تشخیص اشیا، یادگیری عمیق، YOLOv11، رانندگی خودکار، مجموعه‌داده IADD، نشت داده، "
         "میانگین دقت متوسط", 12, jc="start", start=1, end=1)

    # ================================================ ۶ تا ۸. فهرست‌ها (ج، چ، ح)
    # «فاصله تیتر «فهرست مطالب» نسبت به بالای صفحه ۱۱cm است و نسبت به دو طرف صفحه
    #  کاملاً در وسط قرار می‌گیرد و با فونت ۱۴ سیاه لوتوس نوشته می‌شود. یک سانتی‌متر
    #  پایین‌تر از تیتر، با فونت ۱۴ سیاه لوتوس کلمه‌های «عنوان» و «صفحه» نوشته
    #  می‌شود. فاصله «عنوان» تا سمت راست صفحه ۳cm و فاصله «صفحه» تا سمت چپ صفحه ۲cm»
    for _letter, _title, _instr in (
            (LETTERS[5], "فهرست مطالب", 'TOC \\h \\z \\u'),
            (LETTERS[6], "فهرست شکل‌ها", 'TOC \\h \\z \\t "عنوان شکل,1"'),
            (LETTERS[7], "فهرست جدول‌ها", 'TOC \\h \\z \\t "عنوان جدول,1"')):
        new_section(_letter)
        gap(8.0)                                    # ۳ + ۸ = ۱۱
        para(_title, 14, True, font=TITLE_FONT)
        gap(1.0)
        _p = doc.add_paragraph(); _bidi(_p)
        _p.paragraph_format.space_after = Pt(6)
        _tabs = OxmlElement("w:tabs"); _tb = OxmlElement("w:tab")
        _tb.set(qn("w:val"), "end"); _tb.set(qn("w:pos"), str(int(16 * TWIP)))
        _tabs.append(_tb); _p._p.get_or_add_pPr().append(_tabs)
        for _txt in ("عنوان", "\t", "صفحه"):
            _r = _p.add_run(_txt)
            _r.bold = _txt != "\t"; _r.font.size = Pt(14); _r.font.name = TITLE_FONT
            _rs(_r, 14, _r.bold, TITLE_FONT)
        _jc(_p, "start")
        # خط جداکننده زیر «عنوان» و «صفحه»، مثل خود آیین‌نامه
        _pPr = _p._p.get_or_add_pPr()
        _bd = OxmlElement("w:pBdr"); _bt = OxmlElement("w:bottom")
        _bt.set(qn("w:val"), "single"); _bt.set(qn("w:sz"), "8")
        _bt.set(qn("w:space"), "4"); _bt.set(qn("w:color"), "000000")
        _bd.append(_bt); _pPr.append(_bd)
        _field(doc.add_paragraph(), _instr)

    # ============================================ ۹. علائم و اختصارات (صفحه خ)
    # «در این قسمت لیستی از کلیه علائم و اختصاراتی که در متن به‌کار رفته است درج
    #  می‌گردد. نحوه نگارش آن همانند فهرست است. در سمت چپ علامت و در سمت راست مفهوم
    #  آن درج می‌گردد» (ص ۱۵) - پس جدول تمام عرض متن است، نه وسط صفحه.
    new_section(LETTERS[8])
    gap(8.0)
    para("لیست علائم و اختصارات", 14, True, font=TITLE_FONT)
    gap(1.0)

    ABBR = [
        ("ADAS", "سامانه‌های کمک‌راننده پیشرفته"),
        ("AdamW", "بهینه‌ساز آدام با کاهش وزن تفکیک‌شده"),
        ("AP", "میانگین دقت"),
        ("BDD100K", "مجموعه‌داده رانندگی برکلی"),
        ("C2PSA", "بلوک پیچشی با توجه مکانی موازی"),
        ("C3k2", "بلوک پیچشی متقاطع با هسته کوچک"),
        ("CNN", "شبکه عصبی پیچشی"),
        ("COCO", "مجموعه‌داده اشیای متداول در بستر"),
        ("F1", "امتیاز هماهنگ دقت و فراخوانی"),
        ("FN", "منفی نادرست"),
        ("FP", "مثبت نادرست"),
        ("GPU", "پردازنده گرافیکی"),
        ("IADD", "مجموعه‌داده رانندگی خودکار ایران"),
        ("IoU", "نسبت هم‌پوشانی"),
        ("KITTI", "مجموعه‌داده کیتی"),
        ("mAP", "میانگین دقت متوسط"),
        ("R-CNN", "شبکه عصبی پیچشی مبتنی بر ناحیه"),
        ("SGD", "گرادیان کاهشی تصادفی"),
        ("SPPF", "ادغام هرمی مکانی سریع"),
        ("SSD", "آشکارساز تک‌گذر چندکادره"),
        ("TP", "مثبت درست"),
        ("YOLO", "تو فقط یک‌بار نگاه می‌کنی"),
    ]
    _t = doc.add_table(rows=0, cols=2)
    _t.autofit = False
    _t._tbl.tblPr.append(OxmlElement("w:bidiVisual"))       # ستون ۰ سمت راست
    for _sym, _meaning in ABBR:
        _row = _t.add_row()
        _row.cells[0].width, _row.cells[1].width = Cm(11), Cm(5)
        _c = _row.cells[0].paragraphs[0]; _bidi(_c)
        _c.paragraph_format.line_spacing = 1
        _r = _c.add_run(_meaning); _r.font.size = Pt(12); _rs(_r, 12)
        _jc(_c, "start")
        _c = _row.cells[1].paragraphs[0]
        _c.paragraph_format.line_spacing = 1
        _r = _c.add_run(_sym); _r.font.size = Pt(12); _r.font.name = LAT
        _rs(_r, 12, font=LAT)
        _jc(_c, "left")        # این خانه چپ‌به‌راست است، پس left یعنی واقعاً چپ


    return doc


def _standalone():
    global doc
    doc = Document()
    _st = doc.styles["Normal"]
    _st.font.name = FONT; _st.font.size = Pt(14)
    _st.paragraph_format.line_spacing = 1.5; _st.paragraph_format.space_after = Pt(0)
    build_front(doc)
    doc.save(os.path.join(OUT, "صفحات_فرعی.docx"))
    n = len(doc.sections)

    cover = Document()
    _st = cover.styles["Normal"]
    _st.font.name = FONT; _st.font.size = Pt(14)
    _st.paragraph_format.line_spacing = 1.5; _st.paragraph_format.space_after = Pt(0)
    doc = cover
    _page_setup(cover.sections[0], top=2.0, bottom=1.5)
    _footer_letter(cover.sections[0], "")
    title_block(with_judge=False)
    cover.save(os.path.join(OUT, "روی_جلد.docx"))
    print("front matter sections:", n, "| cover: 1 page")


if __name__ == "__main__":
    _standalone()
