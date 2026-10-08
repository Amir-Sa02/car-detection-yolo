from pathlib import Path
from pptx import Presentation
from docx import Document

root=Path(r'C:\Users\amir2\Desktop\نسخه ارسالی\documents')
prs=Presentation(root/'ارائه_دفاع_نهایی_جلسه.pptx')
print('SLIDES',len(prs.slides))
for i,s in enumerate(prs.slides,1):
 print('\n--- SLIDE',i,'---')
 for j,sh in enumerate(s.shapes):
  if sh.has_text_frame and sh.text.strip():
   print(j, f'({sh.left/914400:.2f},{sh.top/914400:.2f},{sh.width/914400:.2f},{sh.height/914400:.2f})', repr(sh.text[:450]))
  elif sh.shape_type==13:
   print(j,'IMAGE', f'({sh.left/914400:.2f},{sh.top/914400:.2f},{sh.width/914400:.2f},{sh.height/914400:.2f})')

d=Document(root/'پایان‌نامه.docx')
print('\n--- LAST PARAGRAPHS ---')
for j,p in list(enumerate(d.paragraphs))[-65:]:
 if p.text.strip(): print(j,repr(p.text[:500]))
