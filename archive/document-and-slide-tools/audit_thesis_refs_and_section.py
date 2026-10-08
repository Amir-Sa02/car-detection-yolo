from docx import Document
from zipfile import ZipFile
from lxml import etree
from collections import Counter
import re, sys
sys.stdout.reconfigure(encoding='utf-8')
p=r'C:\Users\amir2\Desktop\نسخه ارسالی\documents\پایان‌نامه.docx'
d=Document(p)
print('PARAGRAPHS',len(d.paragraphs),'TABLES',len(d.tables))
for i,par in enumerate(d.paragraphs):
 t=' '.join(par.text.split())
 if re.match(r'^\[(?:[0-9]+|[۰-۹]+)\]',t): print('REF',i,t)

with ZipFile(p) as z:
 root=etree.fromstring(z.read('word/document.xml'))
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
 paras=root.xpath('.//w:body//w:p',namespaces=ns)
 for i,par in enumerate(paras):
  txt=''.join(par.xpath('.//w:t/text()',namespaces=ns))
  if txt.strip()=='۱‏-‏۴‏-‏۴- دقت متوسط و میانگین دقت متوسط':
   for j in range(i,i+20):
    q=paras[j]
    plain=''.join(q.xpath('.//w:t/text()',namespaces=ns)).strip()
    math=''.join(q.xpath('.//m:t/text()',namespaces=ns)).strip()
    h=q.xpath('.//w:hyperlink',namespaces=ns)
    print('SEC',j,'TEXT=',plain,'MATH=',math,'HYP=',len(h))
   break
 print('BOOKMARKS')
 for b in root.xpath('.//w:bookmarkStart',namespaces=ns):
  name=b.get('{%s}name'%ns['w'])
  if name and ('ref' in name.lower() or 'bib' in name.lower()): print(name)
 rels=etree.fromstring(z.read('word/_rels/document.xml.rels'))
 hrels=[]
 for rel in rels:
  if 'hyperlink' in (rel.get('Type') or ''): hrels.append((rel.get('Id'),rel.get('Target'),rel.get('TargetMode')))
 print('HYPERLINK_RELS',len(hrels),hrels[:20])
