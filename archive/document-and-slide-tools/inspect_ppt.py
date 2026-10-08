from pptx import Presentation
import sys

sys.stdout.reconfigure(encoding="utf-8")

p = r"D:\projects\car-detection-yolo\thesis\ارائه\ارائه_دفاع_نهایی_جلسه.pptx"
prs = Presentation(p)
print("slides", len(prs.slides))
for slide_no in (29,):
    print("SLIDE", slide_no)
    for i, sh in enumerate(prs.slides[slide_no - 1].shapes, 1):
        t = getattr(sh, "text", "").replace("\n", " | ")
        print(i, sh.name, round(sh.left / 914400, 2), round(sh.top / 914400, 2),
              round(sh.width / 914400, 2), round(sh.height / 914400, 2), repr(t[:220]))
