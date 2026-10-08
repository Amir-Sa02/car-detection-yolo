from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

paper = Path(r"D:\projects\car-detection-yolo\thesis\مراجع\IET Image Processing - 2022 - Khosravian - Multi‐domain autonomous driving dataset  Towards enhancing the generalization of.pdf")
thesis = Path(r"C:\Users\amir2\Desktop\نسخه ارسالی\documents\پایان‌نامه.docx")
deck = Path(r"C:\Users\amir2\Desktop\نسخه ارسالی\documents\ارائه-دفاع.pptx")
out = Path(r"C:\Users\amir2\Desktop\cat-claude\iadd_source_extract")
out.mkdir(exist_ok=True)

r = PdfReader(str(paper))
parts = []
for i, page in enumerate(r.pages, 1):
    text = page.extract_text() or ""
    (out / f"paper_page_{i:02d}.txt").write_text(text, encoding="utf-8")
    parts.append(f"\n===== PDF PAGE {i} =====\n{text}")
(out / "paper_all.txt").write_text("\n".join(parts), encoding="utf-8")

d = Document(str(thesis))
tp = []
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if any(k in t for k in ["IADD", "رانندگی خودکار ایران", "نیمه", "۹۷٬۵۲۸", "۸۳٬۰۴۲", "۱۷۱", "YOLOv4"]):
        tp.append(f"P{i}: {t}")
(out / "thesis_iadd.txt").write_text("\n\n".join(tp), encoding="utf-8")

prs = Presentation(str(deck))
sp = []
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for sh in slide.shapes:
        if hasattr(sh, "text") and sh.text.strip():
            texts.append(sh.text.strip())
    joined = "\n".join(texts)
    if any(k in joined for k in ["IADD", "۹۷٬۵۲۸", "۸۳٬۰۴۲", "۱۷۱", "YOLOv4"]):
        sp.append(f"===== SLIDE {i} =====\n{joined}")
(out / "deck_iadd.txt").write_text("\n\n".join(sp), encoding="utf-8")
print(f"pages={len(r.pages)} thesis_hits={len(tp)} slide_hits={len(sp)} out={out}")
