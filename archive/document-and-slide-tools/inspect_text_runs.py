from pathlib import Path
from pptx import Presentation
p=Presentation(Path(r'C:\Users\amir2\Desktop\نسخه ارسالی\documents\ارائه_دفاع_نهایی_جلسه.pptx'))
for n,indices in {3:[3],6:[],7:[],8:[],9:[4],11:[],14:[4,5,6,7,8,9,10,11,12],16:[5],21:[4],29:[3],30:[0,2,4,6,8,10,12,14,16,18,20,22,24]}.items():
 print('\nSLIDE',n)
 s=p.slides[n-1]
 if not indices: indices=[i for i,sh in enumerate(s.shapes) if sh.has_text_frame and sh.text.strip()]
 for i in indices:
  sh=s.shapes[i]
  if not sh.has_text_frame:continue
  print('SHAPE',i,sh.name,'font',[(r.text,r.font.name,r.font.size.pt if r.font.size else None) for q in sh.text_frame.paragraphs for r in q.runs][:20])
