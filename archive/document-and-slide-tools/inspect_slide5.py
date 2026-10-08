from pptx import Presentation
import sys
sys.stdout.reconfigure(encoding='utf-8')
p=r'D:\projects\car-detection-yolo\thesis\ارائه\ارائه_دفاع_نهایی_جلسه.pptx'
prs=Presentation(p)
for sn in (4,5,6):
 print('SLIDE',sn)
 for i,sh in enumerate(prs.slides[sn-1].shapes,1):
  print(i,sh.name,round(sh.left/914400,2),round(sh.top/914400,2),round(sh.width/914400,2),round(sh.height/914400,2),repr(getattr(sh,'text','').replace('\n',' | ')[:180]))
