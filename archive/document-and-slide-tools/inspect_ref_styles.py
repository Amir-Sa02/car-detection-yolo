from pptx import Presentation
import sys
sys.stdout.reconfigure(encoding="utf-8")
p=r"D:\projects\car-detection-yolo\thesis\ارائه\ارائه_دفاع_نهایی_جلسه.pptx"
s=Presentation(p).slides[26]
for n in (7,8,9,10,11,12,13,14,15,16,17,18,19,20,21,22):
    sh=s.shapes[n-1]
    if not sh.has_text_frame: continue
    for p in sh.text_frame.paragraphs:
      for r in p.runs:
        print(n,sh.name,'font',r.font.name,'size',r.font.size.pt if r.font.size else None,'bold',r.font.bold,'color',getattr(getattr(r.font.color,'_color',None),'rgb',None))
        break
      break
