from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SRC = Path(r"C:\Users\amir2\Desktop\نسخه ارسالی\01_کدهای مطالعه به ترتیب\06__compute_v5_validity.py")
OUT = SRC.parent / "راهنمای_کامل_کد_06_compute_v5_validity.docx"
CHART_DIR = Path(r"D:\projects\car-detection-yolo\docs\dataset_validity")

PERSIAN_FONT = "B Zar"
LATIN_FONT = "Times New Roman"
CODE_FONT = "Consolas"
INK = "111827"
NAVY = "17365D"
PALE = "EAF1F8"
GRID = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = borders.find(qn(f"w:{edge}"))
        if e is None:
            e = OxmlElement(f"w:{edge}")
            borders.append(e)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "5")
        e.set(qn("w:color"), GRID)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def bidi(p, right=True):
    if right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_pr = p._p.get_or_add_pPr()
        b = p_pr.find(qn("w:bidi"))
        if b is None:
            b = OxmlElement("w:bidi")
            p_pr.append(b)
        b.set(qn("w:val"), "1")


LATIN = re.compile(r"[A-Za-z][A-Za-z0-9_.:/@%+\-]*(?:\([^\n)]*\))?")


def add_mixed(p, text, size=14, bold=False, color=INK):
    pos = 0
    for m in LATIN.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            r.font.name = PERSIAN_FONT
            r.font.size = Pt(size)
            r.bold = bold
            r.font.color.rgb = RGBColor.from_string(color)
            r_pr = r._element.get_or_add_rPr()
            rtl = OxmlElement("w:rtl")
            rtl.set(qn("w:val"), "1")
            r_pr.append(rtl)
            rf = r_pr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                r_pr.insert(0, rf)
            rf.set(qn("w:cs"), PERSIAN_FONT)
        r = p.add_run(m.group())
        r.font.name = LATIN_FONT
        r.font.size = Pt(max(9, size - 2))
        r.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        r.font.name = PERSIAN_FONT
        r.font.size = Pt(size)
        r.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        r_pr = r._element.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        r_pr.append(rtl)


def para(doc, text="", style=None, size=14, bold=False, after=5, keep=False):
    p = doc.add_paragraph(style=style)
    bidi(p)
    add_mixed(p, text, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.16
    p.paragraph_format.keep_with_next = keep
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    bidi(p)
    add_mixed(p, text, size=18 if level == 1 else 15, bold=True, color="000000")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph()
    bidi(p)
    p.paragraph_format.right_indent = Cm(0.45 + level * 0.35)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    add_mixed(p, "• " + text, size=13.5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    return p


def code_block(doc, code, caption=None):
    if caption:
        p = doc.add_paragraph()
        bidi(p)
        add_mixed(p, caption, size=12.5, bold=True)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(3)
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F5F7FA")
    set_cell_margins(cell, 120, 150, 120, 150)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(code.splitlines()):
        if i:
            p.add_run("\n")
        r = p.add_run(line)
        r.font.name = CODE_FONT
        r.font.size = Pt(9.3)
        r.font.color.rgb = RGBColor.from_string("1F2937")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def table(doc, headers, rows, widths=None, small=False, ltr_cols=None):
    ltr_cols = set(ltr_cols or [])
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    set_table_borders(t)
    set_repeat_header(t.rows[0])
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        set_cell_shading(c, NAVY)
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        bidi(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_mixed(p, h, size=11.5 if small else 12.5, bold=True, color="FFFFFF")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, value in enumerate(row):
            c = cells[j]
            set_cell_margins(c)
            if i % 2:
                set_cell_shading(c, "F7F9FC")
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            if j in ltr_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(str(value))
                r.font.name = LATIN_FONT
                r.font.size = Pt(9.8 if small else 11)
            else:
                bidi(p)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.RIGHT
                add_mixed(p, str(value), size=10.8 if small else 12)
    if widths:
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def page_break(doc):
    # Let Word paginate naturally. Forced breaks created nearly empty pages
    # when a preceding table flowed onto the next page.
    return None


def add_chart(doc, name, caption):
    path = CHART_DIR / name
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(16.2))
        cp = doc.add_paragraph()
        bidi(cp)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_mixed(cp, caption, size=11.5)
        cp.paragraph_format.space_after = Pt(6)


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = PERSIAN_FONT
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size in (("Title", 22), ("Heading 1", 18), ("Heading 2", 15)):
        s = doc.styles[style_name]
        s.font.name = PERSIAN_FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.element.get_or_add_rPr().set(qn("w:rsidRPr"), "00000000")
        p_pr = s.element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("راهنمای کد اعتبار مجموعه داده  |  صفحه ")
    r.font.name = PERSIAN_FONT
    r.font.size = Pt(9.5)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    return doc


def build():
    doc = setup_doc()
    p = doc.add_paragraph(style="Title")
    bidi(p)
    add_mixed(p, "راهنمای کامل کد محاسبه اعتبار مجموعه داده", size=22, bold=True, color="000000")
    p.paragraph_format.space_after = Pt(5)
    p = doc.add_paragraph()
    bidi(p)
    add_mixed(p, "فایل 06__compute_v5_validity.py", size=15, bold=True, color=NAVY)
    p.paragraph_format.space_after = Pt(10)

    para(doc, "هدف این فایل آن است که بتوانید در جلسه دفاع، منطق کد را از ورودی تا خروجی توضیح دهید؛ بدانید هر عدد جدول و هر ستون نمودار چگونه محاسبه شده است؛ و در برابر پرسش های فنی درباره فرض ها، مرزها و محدودیت های کد پاسخ دقیق بدهید.", size=14)
    para(doc, "نتیجه اصلی: این برنامه مدل را آموزش نمی دهد و مجموعه داده را تغییر نمی دهد. برنامه فقط نسخه نهایی iadd_subset_v5 را می خواند، آمار سه بخش آموزش، اعتبارسنجی و آزمون را محاسبه می کند، نتیجه را در stats.json ذخیره می کند و سه نمودار اولیه می سازد.", size=14, bold=True)

    heading(doc, "خلاصه ای که باید در سی ثانیه بگویید", 1)
    para(doc, "این کد روی هر سه بخش مجموعه داده پیمایش می کند. از نام فایل، شناسه ویدیو و وضعیت محیطی را می گیرد؛ از فایل برچسب YOLO، تعداد نمونه هر رده و مساحت نرمال شده کادرها را می خواند؛ اشیا را به سه گروه کوچک، متوسط و بزرگ تقسیم می کند؛ سپس درصد رده ها، اندازه اشیا و شرایط محیطی را برای هر بخش رسم و همه شمارش ها را در فایل JSON ذخیره می کند.")

    heading(doc, "نقشه کلی اجرای برنامه", 1)
    table(doc, ["گام", "کاری که انجام می شود", "خروجی"], [
        ["۱", "تعریف مسیرها، نام رده ها، آستانه های اندازه و نام سه بخش", "ثابت های برنامه"],
        ["۲", "اجرای تابع analyze برای train و val و test", "سه مجموعه آمار مستقل"],
        ["۳", "ذخیره همه شمارش ها", "stats.json"],
        ["۴", "محاسبه سهم هر رده از نمونه های همان بخش", "class_dist.png"],
        ["۵", "محاسبه سهم سه گروه اندازه", "size_dist.png"],
        ["۶", "محاسبه سهم شرایط محیطی از تصاویر", "cond_dist.png"],
        ["۷", "چاپ خلاصه در خروجی اجرای برنامه", "شمار ویدیو، تصویر، نمونه و اندازه ها"],
    ], widths=[1.2, 10.5, 4.5])

    heading(doc, "ورودی ها و خروجی ها", 1)
    table(doc, ["نوع", "مسیر یا قالب", "معنا"], [
        ["ورودی تصویر", "DS/images/{split}/*.jpg", "فایل های تصویر هر بخش؛ جستجو غیر بازگشتی و فقط برای پسوند jpg کوچک است"],
        ["ورودی برچسب", "DS/labels/{split}/{stem}.txt", "برچسب متناظر هر تصویر در قالب YOLO"],
        ["خروجی آماری", "OUT/stats.json", "شمار تصویر، ویدیو، نمونه، رده، اندازه و وضعیت محیطی"],
        ["خروجی نمودار", "class_dist.png", "درصد نمونه های هر رده در هر بخش"],
        ["خروجی نمودار", "size_dist.png", "درصد اشیای کوچک، متوسط و بزرگ"],
        ["خروجی نمودار", "cond_dist.png", "درصد تصاویر روز، شب، باران و A"],
    ], widths=[2.4, 6.0, 7.8], small=True)
    para(doc, "وابستگی مهم: برنامه فرض می کند نام تصویر چیزی شبیه Record001_D__Record001_D_Record001_D_0.jpg است. بخش پیش از دو زیرخط، یعنی Record001_D، شناسه ویدیو محسوب می شود و حرف پایانی D وضعیت محیطی را نشان می دهد.", size=13.5)

    page_break(doc)
    heading(doc, "توضیح خط های ۱ تا ۱۵", 1)
    code_block(doc, """1  \"\"\"Compute iadd_subset_v5 distribution stats + charts to document its validity.\"\"\"
2  import os, glob, json
3  from collections import Counter
4  import matplotlib
5  matplotlib.use(\"Agg\")
6  import matplotlib.pyplot as plt
8  DS = r\"D:\\projects\\car-detection-yolo\\dataset\\iadd_subset_v5\"
9  OUT = r\"D:\\projects\\car-detection-yolo\\docs\\dataset_validity\"
10 os.makedirs(OUT, exist_ok=True)
11 NAMES = [\"person\", \"car\", \"motorcycle\", \"bus\", \"truck\", \"traffic_light\"]
12 SMALL, LARGE = 0.01, 0.06
13 SPLITS = [\"train\", \"val\", \"test\"]
14 COL = {\"train\": \"#1f4e79\", \"val\": \"#2e75b6\", \"test\": \"#e08a00\"}""", "آماده سازی برنامه")

    table(doc, ["خط", "عنصر", "توضیح دقیق و پاسخ دفاع"], [
        ["۱", "docstring", "هدف فایل را ثبت می کند و روی اجرای برنامه اثری ندارد"],
        ["۲", "os", "کار با مسیر، نام فایل، وجود فایل و ساخت پوشه"],
        ["۲", "glob", "یافتن همه تصاویر jpg در پوشه هر بخش"],
        ["۲", "json", "ذخیره دیکشنری آمار در قالب JSON"],
        ["۳", "Counter", "شمارنده ای که برای کلید دیده نشده مقدار صفر برمی گرداند"],
        ["۴ تا ۶", "matplotlib و Agg", "Agg خروجی را بدون بازکردن پنجره گرافیکی می سازد؛ مناسب سرور و Colab"],
        ["۸", "DS", "ریشه نسخه نهایی مجموعه داده"],
        ["۹ و ۱۰", "OUT", "پوشه خروجی؛ exist_ok=True یعنی اگر پوشه از قبل بود خطا رخ ندهد"],
        ["۱۱", "NAMES", "ترتیب نام ها دقیقا با شناسه های صفر تا پنج در برچسب ها متناظر است"],
        ["۱۲", "SMALL و LARGE", "مرزهای نسبت مساحت کادر به تصویر: یک درصد و شش درصد"],
        ["۱۳", "SPLITS", "ترتیب ثابت اجرای تحلیل و ترتیب میله های نمودار"],
        ["۱۴", "COL", "رنگ هر بخش در هر سه نمودار"],
    ], widths=[1.5, 4.1, 12.0], small=True)

    heading(doc, "چرا از Counter استفاده شده است", 2)
    para(doc, "اگر یک رده در یک بخش هیچ نمونه ای نداشته باشد، دسترسی مستقیم به دیکشنری معمولی ممکن است خطا بدهد. Counter برای کلیدی که هنوز ثبت نشده مقدار صفر می دهد. در انتهای تابع نیز get با مقدار پیش فرض صفر به کار رفته تا خروجی هر شش رده و هر سه گروه اندازه همیشه وجود داشته باشد.")

    heading(doc, "چرا بذر تصادفی وجود ندارد", 2)
    para(doc, "این برنامه هیچ انتخاب تصادفی انجام نمی دهد. برای یک مجموعه فایل ثابت، خروجی آن قطعی است. بذر تصادفی به کد ساخت تقسیم داده مربوط بود، نه به این کد آماری.")

    page_break(doc)
    heading(doc, "تابع analyze و منطق پیمایش داده", 1)
    code_block(doc, """17 def analyze(split):
18     imgs = glob.glob(f\"{DS}/images/{split}/*.jpg\")
19     cls = Counter(); size = Counter(); cond = Counter(); inst = 0; vids = set()
20     for p in imgs:
21         stem = os.path.basename(p)[:-4]
22         rec = stem.split(\"__\")[0]
23         vids.add(rec)
24         cond[rec.split(\"_\")[-1]] += 1
25         t = f\"{DS}/labels/{split}/{stem}.txt\"
26         if os.path.exists(t):
27             for line in open(t):
28                 q = line.split()
29                 if len(q) == 5:
30                     inst += 1; cls[int(q[0])] += 1
31                     a = float(q[3]) * float(q[4])
32                     size[0 if a < SMALL else (2 if a >= LARGE else 1)] += 1
33     return {\"images\": len(imgs), \"videos\": len(vids), \"instances\": inst,
34             \"cls\": {i: cls.get(i, 0) for i in range(6)},
35             \"size\": {k: size.get(k, 0) for k in range(3)},
36             \"cond\": dict(cond)}""", "قلب برنامه")

    table(doc, ["خط", "چه اتفاقی می افتد", "نکته ای که باید بلد باشید"], [
        ["۱۷", "تابع برای یک بخش ورودی تعریف می شود", "ورودی فقط یکی از train، val یا test است"],
        ["۱۸", "فهرست تصاویر jpg همان بخش ساخته می شود", "پوشه های فرعی و JPG بزرگ پیدا نمی شوند"],
        ["۱۹", "شمارنده ها، شمار کل نمونه و مجموعه ویدیوها ساخته می شوند", "set باعث می شود هر ویدیو فقط یک بار شمرده شود"],
        ["۲۰", "تمام تصاویر پیمایش می شوند", "پیچیدگی اصلی برنامه متناسب با تعداد تصویر و سطر برچسب است"],
        ["۲۱", "نام فایل بدون مسیر و پسوند گرفته می شود", "[:-4] فرض می کند پسوند دقیقا چهار نویسه مانند .jpg است"],
        ["۲۲", "بخش پیش از __ به عنوان شناسه ویدیو جدا می شود", "ساختار نام گذاری داده شرط صحت این مرحله است"],
        ["۲۳", "شناسه به مجموعه vids افزوده می شود", "تکرار فریم های یک ویدیو شمار ویدیو را زیاد نمی کند"],
        ["۲۴", "حرف پسوند شناسه ویدیو شمرده می شود", "این شمارش برای هر تصویر است، نه برای هر ویدیو"],
        ["۲۵", "مسیر برچسب متناظر ساخته می شود", "نام پایه تصویر و برچسب باید یکسان باشد"],
        ["۲۶", "فقط اگر فایل برچسب وجود داشته باشد خوانده می شود", "تصویر بدون برچسب همچنان در شمار تصاویر و شرایط هست ولی نمونه ای ندارد"],
        ["۲۷ تا ۲۹", "سطر شکسته و فقط سطر پنج عضوی پذیرفته می شود", "سطر خراب بدون هشدار کنار گذاشته می شود"],
        ["۳۰", "شمار کل نمونه و شمار رده افزایش می یابد", "هر سطر معتبر معادل یک کادر یا instance است"],
        ["۳۱", "عرض و ارتفاع نرمال شده ضرب می شوند", "حاصل، سهم مساحت کادر از مساحت تصویر است"],
        ["۳۲", "نمونه در یکی از سه گروه اندازه قرار می گیرد", "مرزها دقیق و بدون همپوشانی هستند"],
        ["۳۳ تا ۳۶", "دیکشنری نهایی یک بخش برگردانده می شود", "این خروجی مبنای JSON و همه نمودارهاست"],
    ], widths=[1.4, 7.2, 9.0], small=True)

    page_break(doc)
    heading(doc, "نام فایل چگونه تفسیر می شود", 1)
    code_block(doc, """Record001_D__Record001_D_Record001_D_1006.jpg

stem = Record001_D__Record001_D_Record001_D_1006
rec  = Record001_D
condition = D""", "نمونه واقعی از پوشه آموزش")
    bullet(doc, "os.path.basename مسیر پوشه را حذف می کند و فقط نام فایل را نگه می دارد.")
    bullet(doc, "[:-4] پسوند .jpg را حذف می کند.")
    bullet(doc, "split('__')[0] همه چیز پیش از دو زیرخط را به عنوان شناسه ویدیو می گیرد.")
    bullet(doc, "split('_')[-1] آخرین قطعه شناسه، یعنی D را به عنوان کد شرایط می گیرد.")
    para(doc, "پاسخ مهم دفاع: این کد شرایط جوی را از محتوای پیکسل ها تشخیص نمی دهد. کد D، N، R یا A از نام ویدیو استخراج می شود. در نمودار خام، A با برچسب other نمایش داده شده است. نام گذاری آن به عنوان روز ابری در متن نهایی، حاصل بازبینی تصاویر پروژه است و در این تابع انجام نمی شود.", bold=True)

    heading(doc, "قالب برچسب YOLO", 1)
    code_block(doc, """class_id  x_center  y_center  width  height
1         0.500260  0.578704  0.028646  0.040741""", "نمونه یک سطر برچسب")
    table(doc, ["جزء", "اندیس q", "معنا"], [
        ["شناسه رده", "q[0]", "عدد صفر تا پنج"],
        ["مرکز افقی", "q[1]", "مختصات نرمال شده مرکز کادر"],
        ["مرکز عمودی", "q[2]", "مختصات نرمال شده مرکز کادر"],
        ["عرض", "q[3]", "عرض کادر تقسیم بر عرض تصویر"],
        ["ارتفاع", "q[4]", "ارتفاع کادر تقسیم بر ارتفاع تصویر"],
    ], widths=[4.0, 3.0, 9.6])
    para(doc, "چرا برای اندازه فقط q[3] و q[4] لازم است؟ محل مرکز کادر روی اندازه اثر ندارد. چون عرض و ارتفاع بین صفر و یک نرمال شده اند، ضرب آن ها نسبت مساحت کادر به کل تصویر را می دهد. برای نمونه بالا، نسبت مساحت حدود ۰٫۰۰۱۱۷ یا ۰٫۱۱۷ درصد است و در گروه کوچک قرار می گیرد.")

    heading(doc, "تعریف دقیق سه گروه اندازه", 1)
    table(doc, ["گروه", "شرط دقیق کد", "تفسیر"], [
        ["کوچک", "a < 0.01", "مساحت کمتر از یک درصد تصویر"],
        ["متوسط", "0.01 <= a < 0.06", "از یک درصد تا کمتر از شش درصد"],
        ["بزرگ", "a >= 0.06", "شش درصد یا بیشتر"],
    ], widths=[3.0, 5.0, 8.6])
    para(doc, "نکته ظریف: برچسب نمودار خام نوشته است large (>6%)، اما خود کد مقدار دقیقا شش درصد را نیز بزرگ حساب می کند. تعریف عملی و قابل استناد همان شرط کد، یعنی بزرگ تر یا مساوی شش درصد است.", bold=True)

    page_break(doc)
    heading(doc, "ساخت آمار سه بخش و ذخیره فایل JSON", 1)
    code_block(doc, """39 S = {sp: analyze(sp) for sp in SPLITS}
40 json.dump(S, open(f\"{OUT}/stats.json\", \"w\"), indent=2)""", "اجرای تابع برای هر سه بخش")
    para(doc, "دیکشنری سازی خط ۳۹ تابع analyze را سه بار اجرا می کند. در پایان، S سه کلید train، val و test دارد. خط ۴۰ آن را با تورفتگی دو فاصله ذخیره می کند تا فایل برای انسان نیز خوانا باشد.")
    para(doc, "در حافظه پایتون، کلیدهای cls و size عدد صحیح هستند. در استاندارد JSON، کلید شی باید رشته باشد؛ بنابراین در stats.json کلیدهای 0 تا 5 به صورت رشته ذخیره می شوند. به همین دلیل کد نهایی رسم شکل های فارسی از S[sp]['cls'][str(c)] استفاده می کند.")

    heading(doc, "ساختار stats.json", 1)
    code_block(doc, """{
  \"train\": {
    \"images\": 26000,
    \"videos\": 79,
    \"instances\": 204009,
    \"cls\": {\"0\": 15457, \"1\": 173193, ...},
    \"size\": {\"0\": 154015, \"1\": 36954, \"2\": 13040},
    \"cond\": {\"D\": 21416, \"N\": 1977, \"A\": 305, \"R\": 2302}
  },
  \"val\": {...},
  \"test\": {...}
}""", "نمونه خلاصه شده خروجی")

    heading(doc, "فرق images و instances", 1)
    para(doc, "images تعداد فایل های تصویر است. instances تعداد سطرهای معتبر در همه فایل های برچسب و در نتیجه تعداد کادرهای اشیا است. یک تصویر می تواند هیچ کادر، یک کادر یا چندین کادر داشته باشد. عبارت نمونه در تصویر از تقسیم instances بر images به دست می آید؛ برای آموزش برابر ۲۰۴٬۰۰۹ تقسیم بر ۲۶٬۰۰۰ یعنی حدود ۷٫۸ است.")

    page_break(doc)
    heading(doc, "نمودار اول توزیع رده ها", 1)
    code_block(doc, """42 # class distribution
43 fig, ax = plt.subplots(figsize=(8, 4))
44 x = range(6); w = 0.26
45 for i, sp in enumerate(SPLITS):
46     tot = S[sp][\"instances\"]
47     vals = [100 * S[sp][\"cls\"][c] / tot for c in range(6)]
48     ax.bar([xx + (i-1)*w for xx in x], vals, w, label=sp, color=COL[sp])
49 ax.set_xticks(list(x)); ax.set_xticklabels(NAMES, rotation=20)
50 ax.set_ylabel(\"% of split instances\"); ax.set_title(\"Class distribution per split\")
51 ax.legend(); plt.tight_layout(); plt.savefig(f\"{OUT}/class_dist.png\", dpi=140); plt.close()""", "خط های ۴۲ تا ۵۱")
    para(doc, "مخرج کسر، تعداد کل نمونه های همان بخش است. بنابراین هر میله می گوید چند درصد از کادرهای یک بخش متعلق به آن رده است. این نمودار تعداد تصویر را نشان نمی دهد. سه میله کنار هر رده با جابه جایی منفی، صفر و مثبت نسبت به محل اصلی رده رسم می شوند.")
    table(doc, ["رده", "آموزش", "اعتبارسنجی", "آزمون"], [
        ["شخص", "۷٫۵۷۷٪", "۶٫۸۳۸٪", "۶٫۸۹۸٪"],
        ["خودروی سبک", "۸۴٫۸۹۵٪", "۸۳٫۹۷۲٪", "۸۳٫۴۵۱٪"],
        ["موتورسیکلت", "۰٫۸۶۱٪", "۱٫۴۴۳٪", "۱٫۴۸۳٪"],
        ["اتوبوس", "۰٫۵۰۱٪", "۰٫۷۲۰٪", "۰٫۹۲۳٪"],
        ["خودروی باری", "۲٫۵۶۴٪", "۳٫۲۷۶٪", "۳٫۴۴۶٪"],
        ["چراغ راهنمایی", "۳٫۶۰۲٪", "۳٫۷۵۲٪", "۳٫۷۹۸٪"],
    ], widths=[5.0, 3.8, 3.8, 3.8])
    add_chart(doc, "class_dist.png", "خروجی خام همین کد؛ محور عمودی در این فایل خطی است")
    para(doc, "نکته مهم دفاع: این فایل مقیاس لگاریتمی تنظیم نمی کند. شکل فارسی نهایی پایان نامه با فایل 08__make_figures_ch2.py و از روی همین stats.json ساخته شده و در آن محور توزیع رده ها لگاریتمی شده است تا میله های رده های کم تعداد خوانا بمانند.", bold=True)

    page_break(doc)
    heading(doc, "نمودار دوم توزیع اندازه اشیا", 1)
    code_block(doc, """53 # object-size distribution
54 fig, ax = plt.subplots(figsize=(6, 4))
55 bins = [\"small\\n(<1%)\", \"medium\\n(1-6%)\", \"large\\n(>6%)\"]; x = range(3)
56 for i, sp in enumerate(SPLITS):
57     tot = sum(S[sp][\"size\"].values())
58     vals = [100 * S[sp][\"size\"][k] / tot for k in range(3)]
59     ax.bar([xx + (i-1)*w for xx in x], vals, w, label=sp, color=COL[sp])
60 ax.set_xticks(list(x)); ax.set_xticklabels(bins)
61 ax.set_ylabel(\"% of instances\"); ax.set_title(\"Object-size distribution per split\")
62 ax.legend(); plt.tight_layout(); plt.savefig(f\"{OUT}/size_dist.png\", dpi=140); plt.close()""", "خط های ۵۳ تا ۶۲")
    para(doc, "مخرج این درصد، مجموع سه شمارنده اندازه است. در داده سالم این مقدار باید با instances برابر باشد، زیرا هر سطر معتبر برچسب دقیقا در یکی از سه گروه قرار می گیرد.")
    table(doc, ["اندازه", "آموزش", "اعتبارسنجی", "آزمون"], [
        ["کوچک", "۷۵٫۴۹۴٪", "۷۵٫۰۵۳٪", "۷۴٫۷۴۷٪"],
        ["متوسط", "۱۸٫۱۱۴٪", "۱۷٫۲۵۳٪", "۱۸٫۱۶۳٪"],
        ["بزرگ", "۶٫۳۹۲٪", "۷٫۶۹۴٪", "۷٫۰۹۰٪"],
    ], widths=[5.0, 3.8, 3.8, 3.8])
    add_chart(doc, "size_dist.png", "خروجی خام توزیع اندازه اشیا")
    para(doc, "برداشت قابل دفاع: حدود سه چهارم کادرها در هر سه بخش کمتر از یک درصد مساحت تصویر را پوشش می دهند. نزدیکی درصدها نشان می دهد اعتبارسنجی و آزمون از نظر این تعریف اندازه، به آموزش نزدیک اند. این مشاهده به تنهایی ثابت نمی کند دشواری همه بخش ها کاملا برابر است؛ فقط یکی از محورهای بررسی را پوشش می دهد.")

    page_break(doc)
    heading(doc, "نمودار سوم شرایط محیطی", 1)
    code_block(doc, """64 # condition distribution
65 fig, ax = plt.subplots(figsize=(6, 4))
66 conds = [\"D\", \"N\", \"R\", \"A\"]; x = range(4)
67 for i, sp in enumerate(SPLITS):
68     tot = S[sp][\"images\"]
69     vals = [100 * S[sp][\"cond\"].get(c, 0) / tot for c in conds]
70     ax.bar([xx + (i-1)*w for xx in x], vals, w, label=sp, color=COL[sp])
71 ax.set_xticks(list(x)); ax.set_xticklabels([\"day\", \"night\", \"rain\", \"other\"])
72 ax.set_ylabel(\"% of images\"); ax.set_title(\"Condition distribution per split\")
73 ax.legend(); plt.tight_layout(); plt.savefig(f\"{OUT}/cond_dist.png\", dpi=140); plt.close()""", "خط های ۶۴ تا ۷۳")
    para(doc, "برخلاف دو نمودار قبلی، واحد این نمودار تصویر است. چون شرط محیطی ویژگی ویدیو و در نتیجه هر فریم آن است، تعداد تصاویر هر کد بر کل تصاویر همان بخش تقسیم می شود.")
    table(doc, ["شرایط", "آموزش", "اعتبارسنجی", "آزمون"], [
        ["D روز", "۲۱٬۴۱۶  |  ۸۲٫۳۶۹٪", "۴٬۶۵۶  |  ۸۳٫۵۷۶٪", "۳٬۹۹۵  |  ۷۱٫۷۱۱٪"],
        ["N شب", "۱٬۹۷۷  |  ۷٫۶۰۴٪", "۴۵۴  |  ۸٫۱۴۹٪", "۸۵۴  |  ۱۵٫۳۲۹٪"],
        ["R باران", "۲٬۳۰۲  |  ۸٫۸۵۴٪", "۴۳۲  |  ۷٫۷۵۴٪", "۵۴۲  |  ۹٫۷۲۹٪"],
        ["A سایر یا ابری", "۳۰۵  |  ۱٫۱۷۳٪", "۲۹  |  ۰٫۵۲۱٪", "۱۸۰  |  ۳٫۲۳۱٪"],
    ], widths=[4.3, 4.1, 4.1, 4.1], small=True)
    add_chart(doc, "cond_dist.png", "خروجی خام شرایط محیطی؛ A در خود کد other نامیده شده است")
    table(doc, ["شرایط", "تعداد در کل ۳۷٬۱۴۲ تصویر", "سهم کل"], [
        ["روز", "۳۰٬۰۶۷", "۸۱٫۰٪"],
        ["شب", "۳٬۲۸۵", "۸٫۸٪"],
        ["باران", "۳٬۲۷۶", "۸٫۸٪"],
        ["A یا ابری", "۵۱۴", "۱٫۴٪"],
    ], widths=[5.0, 6.0, 5.5])
    para(doc, "این جدول کل در خود فایل 06 چاپ نمی شود، اما مستقیما از جمع cond سه بخش در stats.json به دست می آید. پس اگر داور پرسید عدد ۸۱ درصد از کجا آمده، پاسخ این است: ۲۱٬۴۱۶ + ۴٬۶۵۶ + ۳٬۹۹۵ = ۳۰٬۰۶۷ و سپس ۳۰٬۰۶۷ تقسیم بر ۳۷٬۱۴۲ ضرب در صد.")

    page_break(doc)
    heading(doc, "بخش چاپ خروجی در ترمینال", 1)
    code_block(doc, """75 print(\"=== iadd_subset_v5 stats ===\")
76 for sp in SPLITS:
77     d = S[sp]
78     sz = sum(d[\"size\"].values())
79     print(f\"{sp}: videos={d['videos']} images={d['images']} inst={d['instances']} \"
80           f\"inst/img={d['instances']/d['images']:.1f}  \"
81           f\"size S/M/L={[round(100*d['size'][k]/sz) for k in range(3)]}%\")
82 print(\"charts + stats.json ->\", OUT)""", "خط های ۷۵ تا ۸۲")
    code_block(doc, """=== iadd_subset_v5 stats ===
train: videos=79 images=26000 inst=204009 inst/img=7.8 size S/M/L=[75, 18, 6]%
val: videos=39 images=5571 inst=36810 inst/img=6.6 size S/M/L=[75, 17, 8]%
test: videos=42 images=5571 inst=34994 inst/img=6.3 size S/M/L=[75, 18, 7]%
charts + stats.json -> D:\\projects\\car-detection-yolo\\docs\\dataset_validity""", "خروجی مورد انتظار")
    para(doc, "تابع round در این بخش فقط برای خلاصه چاپی درصد اندازه ها را به عدد صحیح گرد می کند. فایل JSON و نمودارها از شمارش های دقیق استفاده می کنند. مقدار inst/img نیز فقط برای نمایش با یک رقم اعشار قالب بندی شده است.")

    heading(doc, "اعداد نهایی که باید حفظ باشید", 1)
    table(doc, ["بخش", "ویدیو", "تصویر", "نمونه", "نمونه در تصویر"], [
        ["آموزش", "۷۹", "۲۶٬۰۰۰", "۲۰۴٬۰۰۹", "۷٫۸"],
        ["اعتبارسنجی", "۳۹", "۵٬۵۷۱", "۳۶٬۸۱۰", "۶٫۶"],
        ["آزمون", "۴۲", "۵٬۵۷۱", "۳۴٬۹۹۴", "۶٫۳"],
        ["مجموع", "۱۶۰", "۳۷٬۱۴۲", "۲۷۵٬۸۱۳", "۷٫۴"],
    ], widths=[4.0, 2.5, 3.5, 3.5, 3.5])

    heading(doc, "رابطه این فایل با پایان نامه و پاورپوینت", 1)
    bullet(doc, "جدول حجم بخش ها از images، videos، instances و نسبت instances/images ساخته شده است.")
    bullet(doc, "جدول شمار نمونه هر رده از cls استخراج شده است.")
    bullet(doc, "جدول اندازه اشیا از size استخراج شده است.")
    bullet(doc, "جدول شرایط محیطی از cond استخراج شده است.")
    bullet(doc, "سه شکل فارسی فصل دوم مستقیما از stats.json خوانده شده اند، اما ظاهر فارسی و محورهای لگاریتمی در 08__make_figures_ch2.py ساخته شده اند.")

    page_break(doc)
    heading(doc, "محدودیت ها و نکات ظریف کد", 1)
    table(doc, ["موضوع", "رفتار واقعی کد", "پاسخ مناسب در دفاع"], [
        ["تصویر بدون برچسب", "تصویر و شرایط شمرده می شود ولی نمونه ای ندارد", "این کد وجود فایل برچسب را کنترل می کند، اما درباره علت نبود آن هشدار نمی دهد"],
        ["سطر برچسب خراب", "اگر دقیقا پنج جزء نداشته باشد بی صدا رد می شود", "برای ممیزی سخت گیرانه بهتر بود شمار سطرهای ردشده گزارش شود"],
        ["شناسه رده خارج از صفر تا پنج", "inst افزایش می یابد، ولی خروجی cls فقط صفر تا پنج را برمی گرداند", "کد بر سالم بودن شناسه های کلاس تکیه دارد؛ اعتبارسنجی صریح می توانست افزوده شود"],
        ["پسوند تصویر", "فقط *.jpg کوچک", "png، jpeg یا JPG در این نسخه دیده نمی شوند"],
        ["پوشه فرعی", "glob غیر بازگشتی است", "ساختار فعلی داده تخت است؛ برای پوشه فرعی باید recursive فعال شود"],
        ["تقسیم ویدیو", "فقط شمار ویدیو در هر بخش را می گیرد", "این تابع به تنهایی صفر بودن اشتراک ویدیوهای سه بخش را اثبات نمی کند"],
        ["شرایط محیطی", "از پسوند نام ویدیو می خواند", "هیچ مدل بینایی برای تشخیص آب و هوا اجرا نمی شود"],
        ["آستانه اندازه", "قرارداد پروژه است", "این مرزها تعریف عملی تحلیل ما هستند و قانون جهانی YOLO نیستند"],
        ["محور لگاریتمی", "در این فایل وجود ندارد", "نسخه فارسی نهایی در make_figures_ch2.py لگاریتمی شده است"],
        ["تقسیم بر صفر", "برای پوشه خالی محافظ ندارد", "با داده فعلی رخ نمی دهد؛ کد عمومی تر باید این حالت را کنترل کند"],
        ["ترتیب تصاویر", "glob ممکن است ترتیب متفاوتی بدهد", "چون فقط شمارش انجام می شود، ترتیب روی نتیجه اثر ندارد"],
        ["مصرف حافظه", "فقط فهرست مسیرها و شمارنده ها نگه داشته می شوند", "تصویرها باز یا وارد حافظه نمی شوند؛ برنامه سبک است"],
    ], widths=[3.2, 6.5, 7.3], small=True)

    heading(doc, "این کد چه چیزهایی را ثابت نمی کند", 1)
    bullet(doc, "ثابت نمی کند میان سه بخش هیچ ویدیوی مشترکی نیست؛ برای آن باید مجموعه شناسه های سه بخش با هم مقایسه شوند.")
    bullet(doc, "ثابت نمی کند برچسب ها درست هستند؛ فقط سطرهای پنج عضوی را می شمارد.")
    bullet(doc, "ثابت نمی کند سه بخش از همه نظر هم دشواری یکسان دارند؛ فقط سه محور رده، اندازه و شرایط را توصیف می کند.")
    bullet(doc, "هیچ نتیجه ای درباره دقت مدل، mAP یا ماتریس درهم ریختگی تولید نمی کند.")
    bullet(doc, "مجموعه داده را نمی سازد؛ ساخت زیرمجموعه در build_subset_v5.py انجام شده است.")

    page_break(doc)
    heading(doc, "پرسش های احتمالی داور و پاسخ های پیشنهادی", 1)
    qa = [
        ("هدف اصلی این کد چیست؟", "ممیزی آماری نسخه نهایی داده و تولید شمارش های پایه برای جدول ها و نمودارهای فصل دوم است."),
        ("آیا این کد مدل را اجرا می کند؟", "خیر. هیچ وزن، پیش بینی یا معیار مدل در آن وجود ندارد؛ فقط فایل نام و برچسب خوانده می شود."),
        ("چرا تعداد نمونه از تعداد تصویر بیشتر است؟", "هر تصویر می تواند چند شیء داشته باشد و هر سطر برچسب یک نمونه محسوب می شود."),
        ("چطور تعداد ویدیو را حساب کرده اید؟", "شناسه پیش از دو زیرخط از نام هر تصویر استخراج و در set قرار داده شده است؛ طول set تعداد شناسه های یکتا است."),
        ("چطور شرایط محیطی تعیین شده؟", "از آخرین بخش شناسه ویدیو، یعنی D، N، R یا A. کد محتوای تصویر را طبقه بندی نمی کند."),
        ("چرا اندازه را با عرض ضرب در ارتفاع حساب کردید؟", "مختصات YOLO نرمال شده اند؛ پس ضرب عرض و ارتفاع نسبت مساحت کادر به کل تصویر است."),
        ("آستانه یک و شش درصد از کجا آمده؟", "این ها مرزهای عملی تعریف شده برای تحلیل پروژه اند، نه خروجی مدل و نه قانون ذاتی YOLO."),
        ("دقیقا شش درصد در کدام گروه است؟", "به دلیل شرط a >= 0.06 در گروه بزرگ است."),
        ("چرا از مساحت پیکسلی استفاده نکردید؟", "نسبت نرمال شده میان تصاویر با ابعاد متفاوت قابل مقایسه است و به اندازه مطلق تصویر وابسته نیست."),
        ("چرا مرکز کادر در اندازه دخالت ندارد؟", "مرکز فقط محل شیء را مشخص می کند؛ اندازه از عرض و ارتفاع به دست می آید."),
        ("چرا Counter؟", "برای شمارش ساده و بازگرداندن صفر برای کلیدهای دیده نشده."),
        ("چرا set برای ویدیو؟", "چون هزاران فریم از یک ویدیو وجود دارد و شناسه باید فقط یک بار شمرده شود."),
        ("چرا Agg؟", "تا نمودار بدون رابط گرافیکی روی سرور یا محیط headless ذخیره شود."),
        ("چرا سه بخش را جدا تحلیل کردید؟", "برای بررسی نزدیکی توزیع های آموزش، اعتبارسنجی و آزمون و جلوگیری از یک ارزیابی نامتوازن."),
        ("آیا نزدیکی نمودارها برابری کامل دشواری را ثابت می کند؟", "خیر. فقط شواهد توصیفی در سه محور فراهم می کند."),
        ("اگر فایل برچسب نباشد چه می شود؟", "تصویر در شمار تصاویر باقی می ماند، اما هیچ نمونه ای برای آن ثبت نمی شود."),
        ("اگر یک سطر شش مقدار داشته باشد چه می شود؟", "به دلیل len(q) == 5 کنار گذاشته می شود و برنامه هشدار نمی دهد."),
        ("آیا کد نشت داده را بررسی می کند؟", "خیر. شناسه ها را درون هر بخش می شمارد، اما اشتراک setهای سه بخش را کنترل نمی کند."),
        ("اعداد ۸۱ درصد روز از کجا آمده؟", "جمع تصاویر D در سه بخش ۳۰٬۰۶۷ است؛ تقسیم بر ۳۷٬۱۴۲ حدود ۸۱ درصد می شود."),
        ("چرا نمودار نهایی لگاریتمی است ولی این کد نیست؟", "این فایل آمار خام و نمودار اولیه خطی را می سازد. make_figures_ch2.py همان stats.json را می خواند و برای خوانایی رده های کم تعداد محور نهایی را لگاریتمی می کند."),
        ("آیا خروجی قابل بازتولید است؟", "بله، چون هیچ تصادفی وجود ندارد؛ برای فایل های ورودی ثابت، شمارش ثابت است."),
        ("زمان اجرای برنامه به چه چیزی وابسته است؟", "تقریبا به تعداد تصاویر و مجموع سطرهای برچسب؛ خود پیکسل های تصویر خوانده نمی شوند."),
        ("چرا JSON ساخته شده؟", "تا محاسبه آمار از طراحی شکل جدا شود و کدهای دیگر بدون شمارش دوباره داده از نتیجه استفاده کنند."),
        ("کدام فایل شکل های فارسی را می سازد؟", "08__make_figures_ch2.py که stats.json را ورودی می گیرد."),
    ]
    for q, a in qa:
        p = doc.add_paragraph()
        bidi(p)
        add_mixed(p, "پرسش: " + q, size=13.5, bold=True, color=NAVY)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(1)
        para(doc, "پاسخ: " + a, size=13.2, after=5)

    page_break(doc)
    heading(doc, "توضیح نود ثانیه ای پیشنهادی در جلسه", 1)
    para(doc, "این فایل برای اعتبارسنجی آماری زیرمجموعه نهایی نوشته شده است. ابتدا برای هر یک از بخش های آموزش، اعتبارسنجی و آزمون، همه تصاویر jpg را پیدا می کند. از نام هر تصویر، شناسه ویدیو و کد شرایط محیطی استخراج می شود. سپس فایل برچسب متناظر در قالب YOLO خوانده می شود؛ هر سطر پنج عضوی یک نمونه است. شناسه رده از ستون اول گرفته می شود و با ضرب عرض و ارتفاع نرمال شده، سهم مساحت کادر از تصویر به دست می آید. بر اساس مرزهای یک و شش درصد، هر شیء در گروه کوچک، متوسط یا بزرگ قرار می گیرد. خروجی این مرحله برای هر بخش شامل تعداد تصویر، ویدیو، نمونه، نمونه های هر رده، اندازه ها و شرایط محیطی است. همه آمار در stats.json ذخیره می شود. بعد سه نمودار از سهم رده ها، اندازه اشیا و شرایط محیطی ساخته می شود. جدول ها و شکل های فارسی پایان نامه از همین آمار ساخته شده اند. این کد مدل را ارزیابی نمی کند و به تنهایی نشت داده یا درستی برچسب ها را اثبات نمی کند؛ نقش آن توصیف و مقایسه توزیع سه بخش است.")

    heading(doc, "چک فهرست مطالعه", 1)
    for item in [
        "بتوانم فرق image، video و instance را توضیح دهم.",
        "بتوانم از روی نام فایل، rec و condition را استخراج کنم.",
        "قالب پنج عضوی YOLO و نقش q[0] تا q[4] را بدانم.",
        "مرزهای دقیق کوچک، متوسط و بزرگ را حفظ باشم.",
        "بدانم مخرج درصد هر یک از سه نمودار چیست.",
        "اعداد ۷۹، ۳۹، ۴۲ و ۲۶٬۰۰۰، ۵٬۵۷۱، ۵٬۵۷۱ را بدانم.",
        "بتوانم توضیح دهم چرا A در کد other است ولی در متن پروژه ابری نامیده شده است.",
        "بدانم نمودارهای فارسی نهایی را make_figures_ch2.py از stats.json می سازد.",
        "محدودیت های مهم کد را بدون زیر سوال بردن نتیجه توضیح دهم.",
    ]:
        bullet(doc, "□ " + item)

    heading(doc, "مسیر فایل های مرتبط", 1)
    table(doc, ["فایل", "نقش"], [
        [str(SRC), "کد موضوع این راهنما"],
        [r"D:\projects\car-detection-yolo\dataset\iadd_subset_v5", "مجموعه داده ورودی"],
        [r"D:\projects\car-detection-yolo\docs\dataset_validity\stats.json", "آمار خام خروجی"],
        [r"C:\Users\amir2\Desktop\نسخه ارسالی\01_کدهای مطالعه به ترتیب\08__make_figures_ch2.py", "ساخت شکل های فارسی نهایی فصل دوم"],
        [r"D:\projects\car-detection-yolo\colab\build_subset_v5.py", "ساخت خود زیرمجموعه؛ وظیفه ای جدا از این کد"],
    ], widths=[10.3, 6.3], small=True, ltr_cols={0})

    heading(doc, "پیوست کد کامل با شماره خط", 1)
    source_lines = SRC.read_text(encoding="utf-8").splitlines()
    numbered = [f"{i:02d}  {line}" for i, line in enumerate(source_lines, 1)]
    for start in range(0, len(numbered), 28):
        end = min(start + 28, len(numbered))
        code_block(doc, "\n".join(numbered[start:end]), f"خط های {start + 1} تا {end}")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
