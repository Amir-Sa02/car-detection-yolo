from docx import Document
from pathlib import Path
import re
p=Path(r'C:\Users\amir2\Desktop\نسخه ارسالی\مطالعه مقاله IADD\راهنمای_فارسی_مطالعه_مقاله_IADD.docx')
d=Document(p)
paras=list(d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            paras.extend(c.paragraphs)
latin_bad=[]; fa_bad=[]
for para in paras:
    for run in para.runs:
        txt=run.text
        if not txt.strip(): continue
        font=run.font.name
        if re.search(r'[A-Za-z]',txt) and font!='Times New Roman': latin_bad.append((txt,font))
        if re.search(r'[\u0600-\u06FF]',txt) and font!='B Zar': fa_bad.append((txt,font))
print('paragraphs',len(paras),'tables',len(d.tables))
print('latin_bad',len(latin_bad),latin_bad[:3])
print('fa_bad',len(fa_bad),fa_bad[:3])
