from pptx import Presentation
import re, sys
sys.stdout.reconfigure(encoding='utf-8')
p=r'D:\projects\car-detection-yolo\thesis\ارائه\ارائه_دفاع_نهایی_جلسه.pptx'
prs=Presentation(p)
for i,s in enumerate(prs.slides,1):
 t=' '.join(getattr(sh,'text','') for sh in s.shapes)
 nums=sorted(set(re.findall(r'[\[\[]([۰-۹0-9]+)[،,]?\s*(?:([۰-۹0-9]+))?[\]\]]',t)))
 if nums: print(i, nums)
