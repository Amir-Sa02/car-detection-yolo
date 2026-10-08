from pathlib import Path
from zipfile import ZipFile
import os
import tempfile

from pptx import Presentation
from lxml import etree

p = Path(r"C:\Users\amir2\Desktop\نسخه ارسالی\documents\ارائه_دفاع_نهایی_جلسه.pptx")
prs = Presentation(p)
slide = prs.slides[8]
shape = slide.shapes[3]
assert shape.text == "مقایسه رده‌های مصوب و نهایی"
runs = [r for par in shape.text_frame.paragraphs for r in par.runs]
assert runs
runs[0].text = "معماری YOLOv11"
for r in runs[1:]:
    r.text = ""
data = etree.tostring(slide.part._element, xml_declaration=True, encoding="UTF-8", standalone=True)
fd, tmp = tempfile.mkstemp(prefix="slide9_title_", suffix=".pptx", dir=p.parent)
os.close(fd)
try:
    with ZipFile(p) as src, ZipFile(tmp, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, data if info.filename == "ppt/slides/slide9.xml" else src.read(info.filename))
    with ZipFile(tmp) as check:
        assert check.testzip() is None
    os.replace(tmp, p)
except BaseException:
    Path(tmp).unlink(missing_ok=True)
    raise
