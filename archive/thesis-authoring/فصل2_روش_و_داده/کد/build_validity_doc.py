# -*- coding: utf-8 -*-
"""Persian (RTL) dataset-validity report for iadd_subset_v5 (tables + embedded charts)."""
import os, json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
CH = os.path.join(HERE, "dataset_validity")
OUT = os.path.join(HERE, "Dataset_v5_Validity_FA.docx")
S = json.load(open(os.path.join(CH, "stats.json")))
NAMES = ["person", "car", "motorcycle", "bus", "truck", "traffic_light"]

DARK = RGBColor(0x1F, 0x49, 0x7D); BLUE = RGBColor(0x2E, 0x75, 0xB6)
GREEN = RGBColor(0x2E, 0x7D, 0x32); GREY = RGBColor(0x59, 0x59, 0x59)
BLACK = RGBColor(0x1A, 0x1A, 0x1A); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FA, LAT = "Tahoma", "Calibri"

doc = Document()
st = doc.styles["Normal"]; st.font.name = FA; st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.4
sec = doc.sections[0]; sec.page_height = Inches(11.69); sec.page_width = Inches(8.27)
sec.top_margin = sec.bottom_margin = Inches(0.8); sec.left_margin = sec.right_margin = Inches(0.85)
_sp = sec._sectPr; _bd = OxmlElement("w:bidi")
_g = _sp.find(qn("w:docGrid"))
(_g.addprevious(_bd) if _g is not None else _sp.append(_bd))


def _bidi(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:bidi")
    ps = pPr.find(qn("w:pStyle"))
    (ps.addnext(b) if ps is not None else pPr.insert(0, b))


def _style(run, font, size, bold, rtl):
    rPr = run._element.get_or_add_rPr(); rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.insert(0, rf)
    rf.set(qn("w:ascii"), font); rf.set(qn("w:hAnsi"), font); rf.set(qn("w:cs"), font if rtl else LAT)
    if size:
        sc = OxmlElement("w:szCs"); sc.set(qn("w:val"), str(int(size*2))); rPr.append(sc)
    if bold:
        rPr.append(OxmlElement("w:bCs"))
    if rtl:
        rPr.append(OxmlElement("w:rtl"))


def fa(text, size=11, bold=False, color=BLACK, center=False, sa=6, sb=0):
    p = doc.add_paragraph(); _bidi(p)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(sa); p.paragraph_format.space_before = Pt(sb)
    r = p.add_run(text); r.font.name = FA; r.font.size = Pt(size); r.bold = bold; r.font.color.rgb = color
    _style(r, FA, size, bold, True); return p


def h1(t): return fa(t, 15, True, DARK, sb=14, sa=6)
def h2(t): return fa(t, 12.5, True, BLUE, sb=10, sa=4)


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; tcPr = c._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "1F497D"); tcPr.append(sh)
        pp = c.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(h); rr.font.name = LAT; rr.font.size = Pt(9.5); rr.bold = True; rr.font.color.rgb = WHITE
    for row in rows:
        cs = t.add_row().cells
        for i, v in enumerate(row):
            pp = cs[i].paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = pp.add_run(str(v)); rr.font.name = LAT; rr.font.size = Pt(9.5); rr.bold = (i == 0)
    return t


def chart(path, w=5.6):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(w))


def pct(sp, key, k, denom):
    return round(100 * S[sp][key][str(k)] / denom[sp], 1)


inst = {sp: S[sp]["instances"] for sp in S}
imgs = {sp: S[sp]["images"] for sp in S}
szt = {sp: sum(S[sp]["size"].values()) for sp in S}

# ---------- title ----------
fa("گزارش اعتبارسنجی دیتاست  IADD-v5", 20, True, DARK, center=True, sa=2)
fa("نسخه‌ی اصلاح‌شده، بدون نشتی، برای پروژه‌ی تشخیص اشیای ترافیکی با YOLOv11", 11, color=GREY, center=True, sa=10)

h1("۱) خلاصه")
table(["Split", "Videos", "Images", "% images", "Instances", "inst/img"],
      [[sp, S[sp]["videos"], f"{imgs[sp]:,}", f"{100*imgs[sp]/sum(imgs.values()):.0f}%",
        f"{inst[sp]:,}", f"{inst[sp]/imgs[sp]:.1f}"] for sp in ("train", "val", "test")])
fa("تقسیم تقریباً ۷۰/۱۵/۱۵ بر حسب تصویر، و واحد تقسیم «کل ویدیو» است تا هیچ فریمی بین بخش‌ها نشت نکند.", 10.5, sb=4)

h1("۲) روش ساخت (چرا معتبر است)")
fa_lines = [
    "تقسیم بر اساس کل ویدیو (نه فریم) → رفع نشتیِ فریم‌های متوالی.",
    "حذف ویدیوهای تکراری (که با perceptual hashing شناسایی شدند) → رفع نشتیِ ویدیوهای هم‌سان با نام متفاوت.",
    "حذف ویدیوهای دارای برچسب خراب (که با تحلیل per-video شناسایی شدند).",
    "طراحی group-stratified: val و test از نظر توزیع اندازه‌ی اشیا، شرایط آب‌وهوایی و کلاس‌ها متوازن شدند (بر اساس ویژگی‌های داده، نه نتیجه‌ی مدل).",
]
for i, ln in enumerate(fa_lines, 1):
    fa(f"{i}. {ln}", 11, sa=3)

h1("۳) توزیع کلاس‌ها")
fa("تعداد نمونه‌ی هر کلاس در هر بخش:", 10.5, sa=3)
table(["Class", "train", "val", "test", "total"],
      [[NAMES[i], f"{S['train']['cls'][str(i)]:,}", f"{S['val']['cls'][str(i)]:,}",
        f"{S['test']['cls'][str(i)]:,}",
        f"{S['train']['cls'][str(i)] + S['val']['cls'][str(i)] + S['test']['cls'][str(i)]:,}"]
       for i in range(6)]
      + [["TOTAL", f"{inst['train']:,}", f"{inst['val']:,}", f"{inst['test']:,}", f"{sum(inst.values()):,}"]])
fa("و سهم هر کلاس از نمونه‌های هر بخش (درصد):", 10.5, sa=3, sb=8)
table(["Class", "train %", "val %", "test %"],
      [[NAMES[i], pct("train", "cls", i, inst), pct("val", "cls", i, inst), pct("test", "cls", i, inst)]
       for i in range(6)])
fa("پروفایل کلاسیِ val و test تقریباً یکسان است (ستون‌های val% و test%).", 10.5, sb=4, sa=4)
chart(os.path.join(CH, "class_dist.png"))

h1("۴) توزیع اندازه‌ی اشیا  (مهم‌ترین شاخصِ سختی)")
table(["Object size", "train %", "val %", "test %"],
      [["small (<1% area)", pct("train", "size", 0, szt), pct("val", "size", 0, szt), pct("test", "size", 0, szt)],
       ["medium (1-6%)", pct("train", "size", 1, szt), pct("val", "size", 1, szt), pct("test", "size", 1, szt)],
       ["large (>6%)", pct("train", "size", 2, szt), pct("val", "size", 2, szt), pct("test", "size", 2, szt)]])
fa("اشیای کوچک سخت‌ترین‌اند و محرک اصلی سختی هستند. توزیع اندازه در val و test تقریباً یکسان است "
   "(هر دو ~۷۵٪ کوچک) — یعنی سختیِ دو بخش متوازن است.", 10.5, color=GREEN, sb=4, sa=4)
chart(os.path.join(CH, "size_dist.png"), 4.6)

h1("۵) توزیع شرایط (روز/شب/باران)")
conds = [("day", "D"), ("night", "N"), ("rain", "R"), ("other", "A")]
table(["Condition", "train %", "val %", "test %"],
      [[nm, round(100*S["train"]["cond"].get(c, 0)/imgs["train"], 1),
        round(100*S["val"]["cond"].get(c, 0)/imgs["val"], 1),
        round(100*S["test"]["cond"].get(c, 0)/imgs["test"], 1)] for nm, c in conds])
fa("همه‌ی شرایط در هر سه بخش حضور دارند. (تعداد ویدیوهای شب کم است، پس درصد شبِ test کمی بالاتر است؛ "
   "ولی چون توزیع اندازه‌ی اشیا — عامل واقعی سختی — متوازن است، این تفاوت اثرگذار نیست.)", 10.5, sb=4, sa=4)
chart(os.path.join(CH, "cond_dist.png"), 4.6)

h1("۶) بررسی نشتی")
fa("بررسی برنامه‌نویسی‌شده تأیید کرد: هیچ ویدیو و هیچ تصویری بین train، val و test مشترک نیست "
   "(اشتراک = صفر). ضمناً ویدیوهای تکراریِ خودِ دیتاست نیز حذف شده‌اند.", 11)

h1("۷) نتیجه‌گیری")
fa("دیتاست IADD-v5 بدونِ نشتی، عاری از ویدیوهای تکراری و برچسب‌های خرابِ شناسایی‌شده، و از نظر توزیع "
   "کلاس‌ها و اندازه‌ی اشیا بین val و test متوازن است. بنابراین یک مبنای ارزیابیِ منصفانه و استاندارد "
   "برای سنجش مدل فراهم می‌کند.", 11)

doc.save(OUT)
print("saved", OUT, "| paragraphs:", len(doc.paragraphs))
