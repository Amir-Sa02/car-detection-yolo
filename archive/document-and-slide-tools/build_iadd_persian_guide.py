from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT_DIR = Path(r"C:\Users\amir2\Desktop\نسخه ارسالی\مطالعه مقاله IADD")
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / "راهنمای_فارسی_مطالعه_مقاله_IADD.docx"

FA_FONT = "B Zar"
EN_FONT = "Times New Roman"
INK = "18222E"
NAVY = "123A63"
BLUE = "2C73A8"
GREEN = "2F7D4A"
ORANGE = "B75A16"
PALE_BLUE = "EAF3F9"
PALE_GREEN = "EAF5EC"
PALE_ORANGE = "FFF1E8"
PALE_GRAY = "F3F5F7"
WHITE = "FFFFFF"

ARABIC_MARKS = re.compile(r"[\u064B-\u065F\u0670\u06D6-\u06ED]")
LATIN = re.compile(r"([A-Za-z][A-Za-z0-9_.@:+/%-]*(?:[ \t]+[A-Za-z][A-Za-z0-9_.@:+/%-]*)*)")


def clean(text):
    return ARABIC_MARKS.sub("", text)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bidi = tbl_pr.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        tbl_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_para_rtl(paragraph, rtl=True):
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1" if rtl else "0")


def set_run_props(run, *, latin=False, size=14, bold=False, color=INK, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size - 1 if latin else size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = EN_FONT if latin else FA_FONT
    r_pr = run._r.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    font = EN_FONT if latin else FA_FONT
    fonts.set(qn("w:ascii"), font)
    fonts.set(qn("w:hAnsi"), font)
    fonts.set(qn("w:eastAsia"), font)
    fonts.set(qn("w:cs"), font)
    rtl = r_pr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        r_pr.append(rtl)
    rtl.set(qn("w:val"), "0" if latin else "1")


def add_mixed(paragraph, text, *, size=14, bold=False, color=INK, italic=False):
    text = clean(text)
    for part in LATIN.split(text):
        if not part:
            continue
        latin = bool(LATIN.fullmatch(part))
        run = paragraph.add_run(part)
        set_run_props(run, latin=latin, size=size, bold=bold, color=color, italic=italic)
    return paragraph


def add_para(doc, text="", *, size=14, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             before=0, after=5, line=1.2, style=None, keep=False):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    set_para_rtl(p)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    add_mixed(p, text, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 20, 2: 17, 3: 15}
    colors = {1: NAVY, 2: BLUE, 3: GREEN}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para_rtl(p)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    add_mixed(p, text, size=sizes[level], bold=True, color=colors[level])
    if level == 1:
        p_pr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "14")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), BLUE)
        pbdr.append(bottom)
        p_pr.append(pbdr)
    return p


def add_bullet(doc, text, *, color=INK, level=0, size=14):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para_rtl(p)
    p.paragraph_format.left_indent = Cm(0.3 + 0.35 * level)
    p.paragraph_format.right_indent = Cm(0.25 + 0.35 * level)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    add_mixed(p, "• " + text, size=size, color=color)
    return p


def add_callout(doc, title, text, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=180, bottom=130, end=180)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para_rtl(p)
    add_mixed(p, title + "\n", size=14, bold=True, color=accent)
    add_mixed(p, text, size=13.5, color=INK)
    table.rows[0].height = None
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_table(doc, headers, rows, widths=None, font_size=12.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_rtl(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para_rtl(p)
        add_mixed(p, text, size=font_size, bold=True, color=WHITE)
        if widths:
            cell.width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        fill = WHITE if ridx % 2 == 0 else PALE_GRAY
        for i, text in enumerate(row):
            cell = cells[i]
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_para_rtl(p)
            add_mixed(p, str(text), size=font_size, color=INK)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2)
section.bottom_margin = Cm(1.8)
section.right_margin = Cm(2.1)
section.left_margin = Cm(2.1)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

# Default style and language.
normal = doc.styles["Normal"]
normal.font.name = FA_FONT
normal.font.size = Pt(14)
normal._element.rPr.rFonts.set(qn("w:cs"), FA_FONT)
normal._element.rPr.rFonts.set(qn("w:ascii"), FA_FONT)

# Footer page number.
footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para_rtl(p)
add_mixed(p, "صفحه ", size=11, color="66707A")
fld_begin = OxmlElement("w:fldChar")
fld_begin.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = " PAGE "
fld_sep = OxmlElement("w:fldChar")
fld_sep.set(qn("w:fldCharType"), "separate")
run_text = OxmlElement("w:t")
run_text.text = "1"
fld_end = OxmlElement("w:fldChar")
fld_end.set(qn("w:fldCharType"), "end")
r = OxmlElement("w:r")
r_pr = OxmlElement("w:rPr")
r_fonts = OxmlElement("w:rFonts")
r_fonts.set(qn("w:ascii"), EN_FONT)
r_fonts.set(qn("w:hAnsi"), EN_FONT)
r_pr.append(r_fonts)
r.append(r_pr)
r.extend([fld_begin, instr, fld_sep, run_text, fld_end])
p._p.append(r)

# Cover.
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para_rtl(p)
add_mixed(p, "راهنمای فارسی مطالعه مقاله اصلی IADD", size=25, bold=True, color=NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para_rtl(p)
add_mixed(p, "خلاصه تحلیلی، نکات دفاع، ارتباط با پایان نامه و موارد نیازمند احتیاط", size=16, color=BLUE)

doc.add_paragraph()
add_callout(
    doc,
    "مقاله مورد مطالعه",
    "Multi-domain autonomous driving dataset: Towards enhancing the generalization of the convolutional neural networks in new environments\n"
    "Amir Khosravian, Abdollah Amirkhani, Masoud Masih-Tehrani, Alireza Yazdanijoo\n"
    "IET Image Processing، جلد ۱۷، سال ۲۰۲۳، صفحات ۱۲۵۳ تا ۱۲۶۶، DOI: 10.1049/ipr2.12710",
    fill=PALE_BLUE,
    accent=NAVY,
)
add_para(doc, "تهیه شده برای مرور پیش از جلسه دفاع پروژه کارشناسی", size=15, bold=True, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER, before=18)
add_para(doc, "نسخه همراه: مقاله PDF با هایلایت رنگی و یادداشت موضوع هر بخش", size=12.5, color="5E6872", align=WD_ALIGN_PARAGRAPH.CENTER)

page_break(doc)
add_heading(doc, "راهنمای استفاده از این جزوه", 1)
add_para(doc, "این جزوه ترجمه تحت اللفظی مقاله نیست. مطالب مقاله به زبان ساده و متناسب با پروژه شما بازنویسی شده اند تا بتوانید در جلسه دفاع، منشأ هر ادعا را دقیق توضیح دهید.")
add_table(doc, ["رنگ در PDF", "معنا", "نحوه مطالعه"], [
    ["زرد", "مستقیما در پایان نامه یا ارائه استفاده شده", "این بخش ها را با اولویت نخست بخوانید"],
    ["سبز", "دانش تکمیلی مهم برای دفاع", "برای پاسخ به سوال های مفهومی و فنی"],
    ["نارنجی", "نیازمند احتیاط یا متفاوت با وضعیت پروژه", "عبارت مقاله را عینا به پروژه تعمیم ندهید"],
], widths=[3.2, 6.4, 6.7])
add_callout(doc, "قاعده مهم", "هرجا در این جزوه عبارت «یافته پروژه» آمده است، آن مطلب از بررسی فایل های دانلودشده، کدها یا آزمایش های شما به دست آمده و نباید به مقاله IADD نسبت داده شود.", fill=PALE_ORANGE, accent=ORANGE)
add_heading(doc, "ترتیب پیشنهادی مطالعه", 2)
for item in [
    "ابتدا بخش «خلاصه دو دقیقه ای» را بخوانید و بدون نگاه کردن بازگو کنید.",
    "سپس بخش های گردآوری داده، برچسب گذاری، آمار و آزمایش های مقاله را همراه PDF هایلایت شده مرور کنید.",
    "بعد جدول «مقاله چه می گوید و پروژه چه یافته است» را چند بار بخوانید؛ بیشترین احتمال سوال داور در همین تفاوت هاست.",
    "در پایان، پاسخ های آماده و متن سه دقیقه ای را با صدای بلند تمرین کنید.",
]:
    add_bullet(doc, item)

page_break(doc)
add_heading(doc, "۱. شناسنامه و اعتبار مقاله", 1)
add_table(doc, ["مشخصه", "اطلاعات"], [
    ["عنوان", "Multi-domain autonomous driving dataset: Towards enhancing the generalization of the convolutional neural networks in new environments"],
    ["نویسندگان", "Amir Khosravian، Abdollah Amirkhani، Masoud Masih-Tehrani و Alireza Yazdanijoo"],
    ["وابستگی دانشگاهی", "دانشکده مهندسی خودرو، دانشگاه علم و صنعت ایران، تهران"],
    ["نشریه", "IET Image Processing، منتشرشده توسط Wiley از طرف Institution of Engineering and Technology"],
    ["نوع مقاله", "مقاله پژوهشی داوری شده و دسترسی آزاد"],
    ["شناسه", "DOI: 10.1049/ipr2.12710"],
    ["صفحات", "۱۲۵۳ تا ۱۲۶۶، جلد ۱۷، سال ۲۰۲۳"],
    ["دسترسی داده", "https://github.com/ahv1373/IADD"],
], widths=[4.2, 12.2])
add_callout(doc, "پاسخ مناسب در دفاع", "این مقاله، مقاله اصلی معرفی مجموعه داده IADD است و توسط پژوهشگران دانشگاه علم و صنعت ایران در نشریه IET Image Processing منتشر شده است. بنابراین منبع پایه برای معرفی مجموعه داده، فرایند گردآوری، برچسب گذاری و آزمایش های اولیه آن محسوب می شود.", fill=PALE_GREEN, accent=GREEN)
add_para(doc, "نکته: برای ادعای دقیق درباره نمایه شدن نشریه در سال مشخص، باید رکورد رسمی Web of Science یا پایگاه دانشگاه بررسی شود. در جلسه دفاع، بیان «مقاله پژوهشی داوری شده در IET Image Processing با DOI معتبر» دقیق و کافی است.", size=12.5, color=ORANGE)

add_heading(doc, "۲. خلاصه دو دقیقه ای مقاله", 1)
add_para(doc, "مسئله اصلی مقاله این است که مدل های بینایی آموزش دیده روی یک شهر یا یک مجموعه داده، هنگام مواجهه با محیط جدید ممکن است افت عملکرد داشته باشند. نویسندگان برای کاهش این مشکل، مجموعه داده IADD را با تصاویر رانندگی واقعی از چند شهر ایران، مسیرهای شهری و برون شهری، شرایط روز، شب و باران و تفکیک پذیری های مختلف ارائه کرده اند. مجموعه داده شش رده شخص، خودروی سبک، موتورسیکلت، اتوبوس، خودروی باری و چراغ راهنمایی را پوشش می دهد.")
add_para(doc, "مقاله گزارش می کند که IADD شامل ۹۷٬۵۲۸ داده برچسب دار از ۱۷۱ ویدیو است. برای برچسب گذاری، حدود ۳۰٬۰۰۰ تصویر ابتدا دستی برچسب خوردند؛ سپس یک مدل YoloR برای تولید برچسب های بقیه تصاویر آموزش داده شد و خروجی ها به وسیله کارشناسان بازبینی و اصلاح شدند. برچسب ها با قالب YOLO ذخیره شده اند.")
add_para(doc, "برای ارزیابی، چهار آشکارساز Faster R-CNN، EfficientDetB4، YOLOv4 و RetinaNet بررسی شدند. مقاله سه سناریو طراحی کرد تا نشان دهد تنوع دامنه IADD می تواند تعمیم مدل به محیط جدید را بهتر کند. همچنین آزمایش جداگانه ای روی تصاویر شبیه سازی شده CARLA انجام شد. نتیجه کلی مقاله این است که آموزش با داده متنوع IADD نسبت به آموزش با داده محدودتر KITTI، در محیط های جدید عملکرد پایدارتری ایجاد می کند.")

page_break(doc)
add_heading(doc, "۳. هدف و مسئله پژوهش", 1)
add_heading(doc, "۳-۱. مسئله تغییر دامنه", 2)
add_para(doc, "مدل تشخیص شی ممکن است روی داده ای شبیه داده آموزشی عملکرد خوبی داشته باشد، اما با تغییر شهر، نور، آب وهوا، نوع دوربین یا شکل خیابان، عملکرد آن کاهش یابد. مقاله این پدیده را در چارچوب تعمیم به محیط های جدید بررسی می کند. هدف IADD فراهم کردن تنوعی است که مدل را با دامنه های تصویری بیشتری روبه رو کند.")
add_callout(doc, "محل در مقاله", "چکیده و مقدمه، صفحه ۱ فایل PDF؛ مشارکت های اصلی مقاله، صفحه ۲.", fill=PALE_BLUE, accent=BLUE)
add_heading(doc, "۳-۲. منظور از چنددامنه", 2)
for item in [
    "چند شهر با بافت شهری متفاوت",
    "مسیرهای شهری، برون شهری، بزرگراه، خیابان، کوچه، میدان و محیط روستایی",
    "شرایط روز، شب و باران",
    "ترافیک روان و متراکم",
    "تفکیک پذیری ها و دوربین های مختلف",
]:
    add_bullet(doc, item)
add_para(doc, "در دفاع، چنددامنه را صرفا به معنی چند وضعیت آب وهوایی تعریف نکنید. دامنه مجموعه ای از ویژگی های محیط، شهر، مسیر، نور، دوربین و ترافیک است.")

add_heading(doc, "۴. گردآوری و ساخت IADD", 1)
add_heading(doc, "۴-۱. محل گردآوری", 2)
add_para(doc, "مقاله نام هشت شهر یا منطقه را ذکر می کند: تهران، البرز، اصفهان، بوشهر، ملایر، اراک، ساوه و رشت. تصاویر از مسیرهای متنوعی مانند بزرگراه، خیابان، کوچه، میدان و محیط روستایی جمع آوری شده اند. این تنوع یکی از دلایل انتخاب IADD برای پروژه شما بود، زیرا شرایط رانندگی ایران را بهتر از مجموعه ای مانند BDD100K بازتاب می دهد.")
add_heading(doc, "۴-۲. رده ها", 2)
add_table(doc, ["نام مقاله", "معادل استفاده شده در پایان نامه", "نکته"], [
    ["person", "شخص", "می تواند فرد داخل یا خارج خودرو را شامل شود؛ عابر دقیقا هم معنی آن نیست"],
    ["car", "خودروی سبک", "خودروهای سواری"],
    ["motorcycle", "موتورسیکلت", "رده مستقل"],
    ["bus", "اتوبوس", "رده مستقل"],
    ["truck", "خودروی باری", "در داده می تواند وانت تا کامیون را در بر گیرد"],
    ["traffic lights", "چراغ راهنمایی", "اشیای کوچک و دوردست در بسیاری از تصاویر"],
], widths=[3.7, 4.7, 8.2])
add_heading(doc, "۴-۳. مشخصات تصویری", 2)
for item in [
    "تفکیک پذیری های ۶۴۰×۴۸۰، ۱۲۸۰×۷۲۰، ۱۹۲۰×۱۰۸۰ و ۳۸۴۰×۲۱۶۰ پیکسل",
    "ویدیوهای ۳۰ و ۶۰ فریم بر ثانیه",
    "استخراج فریم های کلیدی با نرخ نیم فریم بر ثانیه",
    "نگه داشتن بخشی از فریم های تار برای افزایش مقاومت مدل در شرایط واقعی",
    "ثبت داده RGB و IMU در برخی مسیرها با دوربین Intel RealSense D455",
]:
    add_bullet(doc, item)
add_callout(doc, "محل در مقاله", "بخش ۳-۱ و ۳-۲، صفحات ۵ و ۶ فایل PDF.", fill=PALE_BLUE, accent=BLUE)

page_break(doc)
add_heading(doc, "۵. فرایند برچسب گذاری", 1)
add_para(doc, "برچسب گذاری کامل نزدیک به صد هزار تصویر به صورت دستی پرهزینه بود. نویسندگان ابتدا حدود ۳۰٬۰۰۰ تصویر را دستی برچسب زدند. سپس با این داده یک مدل YoloR آموزش دادند تا برای تصاویر باقی مانده کادر و رده پیشنهادی تولید کند. در مرحله آخر، کارشناسان خروجی های مدل را به صورت چشمی بررسی و اصلاح کردند.")
add_table(doc, ["مرحله", "عمل انجام شده", "هدف"], [
    ["۱", "برچسب گذاری دستی حدود ۳۰٬۰۰۰ تصویر", "ساخت داده اولیه قابل اعتماد"],
    ["۲", "آموزش YoloR با تصاویر دستی", "ساخت برچسب زن اولیه"],
    ["۳", "تولید خودکار کادرها برای تصاویر باقی مانده", "کاهش زمان و هزینه"],
    ["۴", "بازبینی و اصلاح خروجی توسط کارشناسان", "کاهش خطاهای برچسب گذاری خودکار"],
], widths=[2, 7.1, 7.5])
add_para(doc, "برچسب هر شی در قالب YOLO با پنج مقدار ذخیره شده است: شناسه رده، مختصات مرکز کادر در راستای افقی و عمودی، عرض و ارتفاع کادر. چهار مقدار هندسی نسبت به اندازه تصویر نرمال می شوند.")
add_callout(doc, "پاسخ دقیق درباره نویز", "مقاله می گوید خروجی های خودکار بازبینی و اصلاح شده اند؛ مقاله ادعا نمی کند که برچسب ها بدون خطا هستند. وجود برچسب های ناقص یا جابه جا در پروژه شما یک یافته تجربی حاصل از بازبینی داده هاست، نه اعتراف صریح مقاله.", fill=PALE_ORANGE, accent=ORANGE)

add_heading(doc, "۶. آمار رسمی و تفاوت با نسخه دانلودی پروژه", 1)
add_table(doc, ["موضوع", "گزارش مقاله", "وضعیت پروژه شما"], [
    ["کل داده برچسب دار", "۹۷٬۵۲۸ داده و ۱۷۱ ویدیو", "در بسته دانلودی، ۸۳٬۰۴۲ فایل برچسب در بخش های train و val یافت شد"],
    ["بخش test رسمی", "در آمار کلی مقاله جزو تقسیم رسمی معرفی شده", "۱۴٬۴۸۶ تصویر داشت، اما فایل برچسب عمومی برای آن یافت نشد"],
    ["مجموع ۸۳٬۰۴۲ و ۱۴٬۴۸۶", "در مقاله این تفکیک توضیح داده نشده", "دقیقا برابر ۹۷٬۵۲۸ است"],
    ["داده بدون برچسب", "جدول مقایسه مقاله حدود ۷٫۵ هزار داده بدون برچسب جداگانه ذکر می کند", "نباید با ۱۴٬۴۸۶ تصویر test بدون فایل برچسب عمومی یکی فرض شود"],
    ["تقسیم رسمی", "۷۸٪ آموزش، ۷٫۲٪ اعتبارسنجی و ۱۴٫۸٪ آزمون", "پروژه تقسیم جدیدی در سطح کل ویدیو ساخت"],
], widths=[4.1, 5.7, 6.8], font_size=11.8)
add_callout(doc, "پاسخ پیشنهادی", "عدد ۹۷٬۵۲۸ از مقاله آمده است. اما فقدان فایل برچسب برای ۱۴٬۴۸۶ تصویر test را خودمان با بررسی ساختار پوشه های نسخه منتشرشده مشاهده کردیم. به همین دلیل آن بخش در ساخت مجموعه داده نهایی قابل استفاده نبود و تصاویر دارای برچسب train و val را در سطح کل ویدیو دوباره به آموزش، اعتبارسنجی و آزمون تقسیم کردیم.", fill=PALE_GREEN, accent=GREEN)

page_break(doc)
add_heading(doc, "۷. آزمایش ها و مدل های مقاله", 1)
add_para(doc, "مقاله چهار آشکارساز را بررسی کرده است: Faster R-CNN، EfficientDetB4، YOLOv4 و RetinaNet. همه مدل ها با وزن های از پیش آموزش دیده روی COCO آغاز شده اند و سپس با داده هدف تنظیم دقیق شده اند. هدف نویسندگان انتخاب بهترین مدل مطلق نبوده است؛ آن ها می خواستند اثر دامنه آموزشی بر تعمیم را در چند خانواده معماری بررسی کنند.")
add_heading(doc, "۷-۱. سه سناریوی اصلی", 2)
add_table(doc, ["سناریو", "داده آموزش", "داده ارزیابی", "پرسش"], [
    ["۱", "KITTI", "IADD", "آیا مدل آموزش دیده در دامنه دیگر به IADD تعمیم می یابد؟"],
    ["۲", "IADD", "IADD", "عملکرد مدل ها هنگام آموزش و ارزیابی روی IADD چقدر است؟"],
    ["۳", "ترکیب KITTI و IADD", "IADD", "آیا افزودن تنوع دامنه عملکرد را بهتر می کند؟"],
], widths=[2.1, 4.1, 4.1, 6.3])
add_heading(doc, "۷-۲. عدم توازن رده ها", 2)
add_para(doc, "مقاله نشان می دهد تعداد نمونه های رده ها یکسان نیست. نویسندگان توضیح می دهند که عملکرد فقط به تعداد نمونه وابسته نیست؛ شکل متمایز شی، تنوع دامنه و دشواری دیداری نیز مؤثرند. این نکته با تحلیل پروژه شما هم سازگار است: کم نمونه بودن یک رده به تنهایی برای پیش بینی دشواری آن کافی نیست.")
add_callout(doc, "محل در مقاله", "بخش انتخاب مدل و جدول های تنظیمات در صفحات ۷ و ۸؛ سناریوها و نتایج در صفحات ۸ تا ۱۲.", fill=PALE_BLUE, accent=BLUE)

add_heading(doc, "۸. معیارها و نتایج", 1)
add_heading(doc, "۸-۱. AP و mAP", 2)
add_para(doc, "مقاله دقت متوسط یا AP را مساحت زیر منحنی دقت برحسب فراخوانی تعریف می کند. سپس میانگین AP رده ها، mAP را تشکیل می دهد. این توضیح، مبنای کلی استفاده از mAP در پایان نامه شماست. با این حال، مقاله در جدول نتیجه YOLOv4 به روشنی مشخص نمی کند mAP گزارش شده در چه آستانه هم پوشانی محاسبه شده است.")
add_heading(doc, "۸-۲. نتیجه سناریوی دوم", 2)
add_table(doc, ["رده", "AP گزارش شده برای YOLOv4"], [
    ["شخص", "۸۶٫۶"],
    ["خودروی سبک", "۹۶٫۱"],
    ["موتورسیکلت", "۸۳٫۴"],
    ["اتوبوس", "۸۱٫۵"],
    ["خودروی باری", "۹۰٫۱"],
    ["چراغ راهنمایی", "۸۱٫۹"],
    ["mAP", "۸۶٫۶"],
], widths=[8.3, 8.3])
add_callout(doc, "چرا با ۰٫۸۷۴ پروژه مقایسه مستقیم نکردیم؟", "زیرا بخش test عمومی IADD برای ما برچسب قابل استفاده نداشت، زیرمجموعه و تقسیم ویدیویی پروژه متفاوت است، معماری و تنظیمات آموزش یکسان نیستند و مقاله آستانه IoU معیار ۸۶٫۶ را در بخش نتیجه به روشنی بیان نکرده است. بنابراین کنار هم گذاشتن ۸۶٫۶ و ۸۷٫۴ می تواند برداشت نادرست از برتری یکی از دو مدل ایجاد کند.", fill=PALE_ORANGE, accent=ORANGE)

page_break(doc)
add_heading(doc, "۹. آزمایش تعمیم با CARLA", 1)
add_para(doc, "مقاله برای بررسی تعمیم به دامنه ای کاملا متفاوت، از تصاویر شبیه سازی شده CARLA استفاده کرده است. مدل هایی که با IADD آموزش دیده بودند، در این داده شبیه سازی شده عموما بهتر از مدل های آموزش دیده با KITTI عمل کردند. برداشت نویسندگان این است که تنوع شهر، مسیر، نور و شرایط محیطی IADD باعث شده مدل نشانه های عمومی تری یاد بگیرد.")
add_bullet(doc, "تعداد کل تصاویر CARLA گزارش شده: ۵٬۱۹۴ تصویر")
add_bullet(doc, "شرایط گزارش شده: روز، شب و باران")
add_bullet(doc, "کاربرد این آزمایش: سنجش تعمیم به دامنه ای خارج از داده واقعی IADD")
add_callout(doc, "نکته احتیاطی", "آزمایش CARLA بخشی از پروژه شما نیست. در دفاع فقط می توانید از آن برای توضیح هدف و اعتبار ادعای تعمیم در مقاله IADD استفاده کنید؛ نتیجه آن را به YOLOv11 آموزش دیده در پروژه خود نسبت ندهید.", fill=PALE_ORANGE, accent=ORANGE)

add_heading(doc, "۱۰. نتیجه گیری و پیشنهادهای آینده مقاله", 1)
add_para(doc, "نتیجه اصلی مقاله این است که افزایش تنوع دامنه آموزشی می تواند توانایی مدل را برای کار در محیط های جدید بهتر کند. نویسندگان IADD را مجموعه ای از تصاویر واقعی رانندگی در ایران معرفی می کنند که از نظر شهر، مسیر، شرایط جوی، نور و تفکیک پذیری تنوع دارد. آزمایش های مقاله نشان می دهند مدل آموزش دیده با IADD در چند سناریوی خارج از دامنه عملکرد مناسبی داشته است.")
add_para(doc, "نویسندگان برای کار آینده پیشنهاد می کنند داده از شهرها و شرایط بیشتری، به ویژه مه، جمع آوری شود و وظایف دیگری مانند بخش بندی معنایی و داده حسگرهایی مانند LIDAR، دوربین استریو و RADAR نیز گسترش یابد.")
add_callout(doc, "محل در مقاله", "بخش Conclusion and Future Works، صفحه ۱۳ فایل PDF.", fill=PALE_BLUE, accent=BLUE)

page_break(doc)
add_heading(doc, "۱۱. ارتباط مقاله با پایان نامه و ارائه شما", 1)
add_table(doc, ["مطلب", "صفحه مقاله", "کاربرد در پایان نامه یا ارائه"], [
    ["هدف IADD و تعمیم", "۱ و ۲", "معرفی مسئله و دلیل انتخاب داده بومی"],
    ["۹۷٬۵۲۸ داده، ۱۷۱ ویدیو و شش رده", "۱، ۲ و ۶", "معرفی مجموعه داده و اسلاید آمار IADD"],
    ["هشت شهر و انواع مسیر", "۲ و ۵", "توضیح تنوع جغرافیایی و محیطی"],
    ["تفکیک پذیری ها و نیم فریم بر ثانیه", "۵", "مشخصات داده و نحوه استخراج تصاویر"],
    ["۳۰٬۰۰۰ برچسب دستی و YoloR", "۵ و ۶", "شرح برچسب گذاری نیمه خودکار و منشأ احتمالی خطا"],
    ["قالب برچسب YOLO", "۶", "توضیح ساختار فایل های برچسب"],
    ["پیش آموزش COCO و یادگیری انتقالی", "۸", "پشتوانه مفهومی راهبرد آموزش"],
    ["تعریف AP و mAP", "۸", "معیارهای ارزیابی"],
    ["نتیجه YOLOv4 برابر ۸۶٫۶", "۹", "آمادگی برای سوال درباره مقایسه با مقاله اصلی"],
    ["آزمایش CARLA", "۱۲ و ۱۳", "توضیح شواهد مقاله درباره تعمیم"],
    ["مخزن GitHub", "۱۳", "لینک دریافت داده و مستندات"],
], widths=[5.2, 2.3, 9.1], font_size=11.7)

add_heading(doc, "۱۲. مقاله چه می گوید و چه نمی گوید", 1)
add_table(doc, ["موضوع", "وضعیت دقیق"], [
    ["نبود برچسب test عمومی", "یافته مستقیم از نسخه دانلودشده پروژه است؛ مقاله آن را صریح توضیح نمی دهد"],
    ["نشت فریم میان بخش ها", "با بررسی ساختار فایل های منتشرشده و شناسه ویدیوها در پروژه کشف شد؛ ادعای مقاله نیست"],
    ["سه ویدیوی تکراری", "یافته پروژه از مقایسه ویدیوهاست"],
    ["دو ویدیوی دارای برچسب معیوب", "یافته پروژه از آمار هم پوشانی و بازبینی چشمی است"],
    ["کدهای D، N، R و A", "مقاله این حروف و معنای A را تعریف نمی کند؛ نگاشت پروژه از نام پوشه ها و بازبینی داده به دست آمده"],
    ["وجود نویز برچسب", "مقاله فرایند بازبینی کارشناسی را ذکر می کند؛ نمونه های نویزی مشخص را پروژه شما کشف کرده است"],
    ["مقایسه ۸۶٫۶ و ۸۷٫۴", "مقایسه مستقیم معتبر نیست؛ تقسیم، داده، مدل، کد و تعریف معیار یکسان نیست"],
    ["تقسیم ویدیویی پروژه", "طراحی خود پروژه برای جلوگیری از نشت و هم تراز کردن ارزیابی است"],
], widths=[5.4, 11.2], font_size=11.7)
add_callout(doc, "مهم ترین جمله دفاعی", "مقاله منبع معرفی IADD و روش ساخت اولیه آن است؛ اما مشکلات نسخه منتشرشده، ساخت زیرمجموعه نهایی، حذف موارد تکراری و معیوب و تقسیم در سطح ویدیو، دستاوردهای خود پروژه ما هستند.", fill=PALE_GREEN, accent=GREEN)

page_break(doc)
add_heading(doc, "۱۳. سوال های محتمل دفاع و پاسخ کوتاه", 1)
qas = [
    ("چرا IADD را انتخاب کردید؟", "زیرا داده واقعی رانندگی در ایران است و از نظر شهر، مسیر، نور و شرایط جوی تنوع دارد. مقاله اصلی نیز هدف مجموعه را بهبود تعمیم در محیط های جدید معرفی می کند."),
    ("این مقاله توسط چه کسانی نوشته شده است؟", "چهار پژوهشگر دانشکده مهندسی خودرو دانشگاه علم و صنعت ایران: Khosravian، Amirkhani، Masih-Tehrani و Yazdanijoo."),
    ("مقاله کجا چاپ شده است؟", "در IET Image Processing، جلد ۱۷، سال ۲۰۲۳، با DOI برابر 10.1049/ipr2.12710."),
    ("IADD چند تصویر و چند ویدیو دارد؟", "مقاله ۹۷٬۵۲۸ داده برچسب دار از ۱۷۱ ویدیو گزارش می کند. نسخه دانلودی مورد استفاده ما ۸۳٬۰۴۲ فایل برچسب در train و val داشت و ۱۴٬۴۸۶ تصویر test فایل برچسب عمومی نداشت."),
    ("چرا همه ۹۷٬۵۲۸ تصویر را استفاده نکردید؟", "چون برای تصاویر بخش test رسمی فایل برچسب در بسته عمومی در دسترس نبود و بدون حقیقت مبنا امکان آموزش و ارزیابی آن ها وجود نداشت."),
    ("برچسب گذاری چگونه انجام شده؟", "حدود ۳۰ هزار تصویر دستی برچسب خورد، YoloR روی آن آموزش دید، بقیه برچسب ها را تولید کرد و کارشناسان خروجی ها را بازبینی و اصلاح کردند."),
    ("اگر کارشناسان اصلاح کردند، چرا نویز دیدید؟", "بازبینی انسانی احتمال خطا را کم می کند اما تضمین صفر خطا نیست. ما موارد مشخصی از کادرهای ناقص، جابه جا یا گمشده را در نسخه منتشرشده مشاهده کردیم."),
    ("چرا رده truck را خودروی باری ترجمه کردید؟", "زیرا نمونه های این رده از وانت تا کامیون را شامل می شوند و «کامیون» دامنه رده را محدود و نادرست نشان می داد."),
    ("چرا نتیجه خود را با YOLOv4 مقاله مقایسه نکردید؟", "داده آزمون، تقسیم، مدل، کد، تنظیمات و حتی آستانه دقیق معیار یکسان نیستند؛ بنابراین مقایسه مستقیم منصفانه و علمی نبود."),
    ("نتیجه YOLOv4 مقاله چه بود؟", "در سناریوی دوم، mAP برابر ۸۶٫۶ گزارش شده است؛ اما آستانه IoU این mAP در جدول و توضیح همان بخش روشن نیست."),
    ("منظور از تعمیم چیست؟", "توانایی مدل برای حفظ عملکرد هنگام مواجهه با شهر، نور، آب وهوا، دوربین یا محیطی که با داده آموزشی دقیقا یکسان نیست."),
    ("CARLA چه نقشی داشت؟", "فقط در مقاله اصلی برای آزمون تعمیم به داده شبیه سازی شده استفاده شد و جزو آزمایش های پروژه ما نبود."),
    ("آیا مقاله تقسیم ویدیویی شما را پیشنهاد کرده؟", "خیر. تقسیم ویدیویی و کنترل نشت، راهکار خود پروژه بر اساس ساختار داده منتشرشده و اصول ارزیابی معتبر بود."),
    ("آیا مقاله کد D/N/R/A را تعریف کرده؟", "خیر. این کدها از نام پوشه های داده منتشرشده استخراج شدند و معنای آن ها در پروژه بررسی شد."),
]
for q, a in qas:
    add_para(doc, q, size=14, bold=True, color=NAVY, before=5, after=2, keep=True)
    add_para(doc, a, size=13.2, after=5)

page_break(doc)
add_heading(doc, "۱۴. متن سه دقیقه ای برای معرفی مقاله در دفاع", 1)
add_callout(doc, "متن پیشنهادی", "مجموعه داده اصلی این پروژه IADD است که توسط پژوهشگران دانشکده مهندسی خودرو دانشگاه علم و صنعت ایران معرفی شده و مقاله آن در نشریه IET Image Processing منتشر شده است. هدف این مجموعه داده، افزایش تنوع دامنه و بهبود تعمیم مدل های تشخیص شی در محیط های جدید است. طبق مقاله، IADD شامل ۹۷٬۵۲۸ داده برچسب دار از ۱۷۱ ویدیو و شش رده شخص، خودروی سبک، موتورسیکلت، اتوبوس، خودروی باری و چراغ راهنمایی است. داده ها در هشت شهر ایران و در مسیرها، نورها، شرایط جوی و تفکیک پذیری های مختلف گردآوری شده اند.\n\nبرای کاهش هزینه برچسب گذاری، حدود ۳۰ هزار تصویر ابتدا دستی برچسب خوردند. سپس مدل YoloR برای برچسب گذاری بقیه تصاویر به کار رفت و خروجی آن توسط کارشناسان بازبینی شد. با این حال، ما در نسخه منتشرشده چند مشکل عملی مشاهده کردیم: فایل برچسب بخش test عمومی در دسترس نبود، فریم های یک ویدیو میان بخش های رسمی پراکنده بودند، چند ویدیوی تکراری وجود داشت و دو ویدیو برچسب های معیوب داشتند. این موارد یافته خود پروژه ما هستند و مقاله آن ها را گزارش نکرده است.\n\nبه همین دلیل، از تصاویر دارای برچسب استفاده کردیم، موارد تکراری و معیوب را کنار گذاشتیم و تقسیم جدید را در سطح کل ویدیو ساختیم تا نشت اطلاعات رخ ندهد. مقاله اصلی برای YOLOv4 مقدار mAP برابر ۸۶٫۶ گزارش کرده است، اما چون داده آزمون، تقسیم، مدل، تنظیمات و آستانه دقیق معیار با پروژه ما یکسان نیست، مقایسه مستقیم آن با نتیجه ۰٫۸۷۴ پروژه علمی و منصفانه نیست.", fill=PALE_BLUE, accent=NAVY)

add_heading(doc, "۱۵. چک لیست مطالعه", 1)
for item in [
    "نام مقاله، نویسندگان، دانشگاه، نشریه و DOI را بدون نگاه کردن بگویید.",
    "عددهای ۹۷٬۵۲۸، ۱۷۱، ۳۰٬۰۰۰، ۸۳٬۰۴۲ و ۱۴٬۴۸۶ را با منشأ هر کدام توضیح دهید.",
    "چهار مرحله برچسب گذاری را به ترتیب بازگو کنید.",
    "شش رده را با ترجمه پایان نامه نام ببرید.",
    "سه سناریوی آزمایش مقاله را توضیح دهید.",
    "دلیل نامعتبر بودن مقایسه مستقیم ۸۶٫۶ با ۸۷٫۴ را در یک دقیقه بیان کنید.",
    "پنج موردی را که مقاله نمی گوید اما پروژه کشف کرده است، نام ببرید.",
    "متن سه دقیقه ای را دو بار با زمان سنج تمرین کنید.",
]:
    add_bullet(doc, "□ " + item)

add_heading(doc, "پیوندهای اصلی", 2)
add_bullet(doc, "صفحه مقاله: https://ietresearch.onlinelibrary.wiley.com/doi/full/10.1049/ipr2.12710")
add_bullet(doc, "مخزن مجموعه داده: https://github.com/ahv1373/IADD")

# Keep all text free of Arabic vowel marks and set document core properties.
doc.core_properties.title = clean("راهنمای فارسی مطالعه مقاله IADD")
doc.core_properties.subject = clean("مطالعه مقاله اصلی مجموعه داده برای دفاع پروژه")
doc.core_properties.author = "محمد امیر صادق زاده"
doc.core_properties.keywords = "IADD, YOLOv11, object detection, thesis defense"

doc.save(OUT)
print(OUT.name.encode("unicode_escape").decode())
