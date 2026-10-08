from pathlib import Path
import json
from pptx import Presentation
from docx import Document

root=Path(r'C:\Users\amir2\Desktop\نسخه ارسالی')
p=Presentation(root/'documents'/'ارائه_دفاع_نهایی_جلسه.pptx')
print('SLIDES', len(p.slides))
for no in (13,14,16,17,18):
 s=p.slides[no-1]
 print('\nSLIDE',no)
 for sh in s.shapes:
  if sh.has_text_frame:
   print(sh.text[:400].replace('\n',' | '))

d=Document(root/'documents'/'پایان‌نامه.docx')
for i,t in enumerate(d.tables):
 rows=[' | '.join(c.text.replace('\n',' ') for c in row.cells) for row in t.rows]
 if i in (12,13):
  print('\nTABLE',i, '\n'+'\n'.join(rows[:15]))

for path in (Path(r'D:\projects\car-detection-yolo\docs\dataset_validity\stats.json'),):
 print('\nSTATS',path.read_text(encoding='utf-8')[:5000])
