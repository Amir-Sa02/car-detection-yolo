from pathlib import Path
from docx import Document


PATH = Path(r"D:\projects\car-detection-yolo\thesis\ارائه\متن_ارائه.docx")
TMP = PATH.with_name("متن_ارائه.tmp.docx")
doc = Document(PATH)


def set_paragraph(paragraph, text):
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def set_cell(cell, text):
    set_paragraph(cell.paragraphs[0], text)
    for paragraph in cell.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)


# Update timing guidance after the section reordering.
old = ("سرعت گفتار معمولی فارسی حدود صد و بیست کلمه در دقیقه است؛ زمان هر اسلاید با همین سرعت تنظیم شده. "
       "اگر عقب افتادی، اسلایدهای ۸، ۹ و ۲۵ را می‌توان کوتاه‌تر گفت. اسلایدهای ۱۳، ۱۵ و ۲۲ را کوتاه نکن؛ "
       "آن‌ها استدلال اصلی پروژه‌اند.")
new = ("سرعت گفتار معمولی فارسی حدود صد و بیست کلمه در دقیقه است و زمان هر اسلاید با همین سرعت تنظیم شده. "
       "اگر عقب افتادی، توضیح اسلایدهای ۸، ۹ و ۱۷ را کوتاه‌تر کن. اسلایدهای ۱۳، ۱۵ و ۲۵ را کوتاه نکن؛ "
       "زیرا به‌ترتیب نشت داده، ممیزی برچسب و نتیجه نهایی پروژه را توضیح می‌دهند.")
matches = [p for p in doc.paragraphs if p.text == old]
if len(matches) != 1:
    raise RuntimeError("timing guidance not found")
set_paragraph(matches[0], new)

# Four-part timing table matching the current 31-slide deck.
t = doc.tables[0]
while len(t.rows) > 6:
    t._tbl.remove(t.rows[-1]._tr)
rows = [
    ["بخش", "اسلایدها", "دقیقه"],
    ["مبانی و صورت مسئله", "۱ تا ۹", "۷٫۵"],
    ["روش پژوهش و آماده‌سازی داده", "۱۰ تا ۲۰", "۱۳"],
    ["نتایج و تحلیل", "۲۱ تا ۲۸", "۸"],
    ["جمع‌بندی و پیشنهادها", "۲۹ تا ۳۱", "۱٫۵"],
    ["جمع", "", "۳۰"],
]
for row, values in zip(t.rows, rows):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value)

# Correct the two contextual Q&A boxes attached to slides 3 and 6.
set_cell(doc.tables[1].cell(0, 0),
         "اگر پرسیدند: چرا از ابتدا مجموعه‌داده‌ای با همان رده‌های مصوب انتخاب نکردید؟\n"
         "هنگام تصویب موضوع، اینترنت قطع بود و امکان بررسی دقیق مجموعه‌داده‌ها وجود نداشت. پس از دسترسی، IADD با نظر استاد راهنما به‌دلیل داشتن صحنه‌های واقعی ایران انتخاب شد. بازبرچسب‌گذاری کامل نیز با زمان و امکانات موجود عملی نبود.")

set_cell(doc.tables[2].cell(0, 0),
         "اگر پرسیدند: چرا IADD را انتخاب کردید؟\n"
         "مدل با وزن‌های ازپیش‌آموخته COCO آغاز شد و سپس با IADD تنظیم دقیق گردید. IADD صحنه‌های واقعی ایران، وسایل نقلیه رایج و شرایط محیطی متنوع را در اختیار پروژه می‌گذاشت و برای مسئله هدف مناسب‌تر بود.")

doc.save(TMP)
check = Document(TMP)
if len(check.tables) != 10:
    raise RuntimeError("table verification failed")
TMP.replace(PATH)
print("speaker tables aligned")
