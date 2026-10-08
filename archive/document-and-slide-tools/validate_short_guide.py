from docx import Document
from pathlib import Path
import re
p=Path(r'C:\Users\amir2\Desktop\نسخه ارسالی\مطالعه مقاله IADD\راهنمای_فارسی_مطالعه_مقاله_IADD.docx')
d=Document(p)
paras=list(d.paragraphs)
for t in d.tables:
  for row in t.rows:
    for c in row.cells: paras.extend(c.paragraphs)
text=''.join(x.text for x in paras)
marks=re.findall(r'[\u064B-\u065F\u0670\u06D6-\u06ED]',text)
latin_bad=[]; fa_bad=[]
for par in paras:
  for run in par.runs:
    if re.search(r'[A-Za-z]',run.text) and run.font.name!='Times New Roman': latin_bad.append(run.text)
    if re.search(r'[\u0600-\u06FF]',run.text) and run.font.name!='B Zar': fa_bad.append(run.text)
print('paragraphs',len(paras),'tables',len(d.tables),'chars',len(text),'marks',len(marks),'latin_bad',len(latin_bad),'fa_bad',len(fa_bad))
