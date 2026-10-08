# -*- coding: utf-8 -*-
"""The single 'مراجع و منابع' section that closes the thesis.

آیین‌نامه: one list, numbered from 1, ordered alphabetically (Persian sources
first, then Latin, then web sources); every entry on one line with the fields
separated by commas and the TITLE IN ITALIC.

Every source in this thesis is in English, so the entries are laid out
left-to-right and left-aligned - the same treatment the guide gives its own
Latin-only footnotes on page 19.
"""
import os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
os.environ["THESIS_OUT"] = os.path.join(HERE, "مراجع_و_منابع.docx")
sys.path.insert(0, os.path.normpath(os.path.join(HERE, "..", "00_قالب_و_آیین‌نامه")))

from docx.enum.text import WD_ALIGN_PARAGRAPH                             # noqa: E402
from docx.shared import Pt, Cm                                            # noqa: E402
from thesis_style import h_ch, doc, finish, LAT, BODY, _rs                # noqa: E402
from bibliography import entries                                          # noqa: E402

h_ch("مراجع و منابع")

for num, authors, title, rest, is_web in entries():
    p = doc.add_paragraph()                    # no bidi: the entry is pure Latin
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    pf.left_indent = Cm(1.0)
    pf.first_line_indent = Cm(-1.0)            # hanging indent lines the text up

    def run(text, italic=False):
        r = p.add_run(text)
        r.font.name = LAT
        r.font.size = Pt(BODY - 2)
        r.italic = italic
        _rs(r, BODY - 2, False, False)
        return r

    run(f"[{num}] ")
    run(authors + ", ")
    run(title, italic=True)                    # آیین‌نامه: title in italic
    if is_web:
        # «آدرس کامل سایتی که اطلاعات فوق را در خود جای داده است، در یک خط
        #  مستقل و از سمت چپ آورده می‌شود» (ص ۱۷)
        run(".")
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.LEFT
        qf = q.paragraph_format
        qf.space_after = Pt(6); qf.line_spacing = 1.5; qf.left_indent = Cm(1.0)
        r = q.add_run(rest); r.font.name = LAT; r.font.size = Pt(BODY - 2)
        _rs(r, BODY - 2, False, False)
    else:
        run(", " + rest + ".")

finish()
