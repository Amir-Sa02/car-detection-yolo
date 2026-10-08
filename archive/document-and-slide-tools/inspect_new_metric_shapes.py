from pptx import Presentation
import sys
sys.stdout.reconfigure(encoding='utf-8')
p=r'D:\projects\car-detection-yolo\thesis\ارائه\ارائه_دفاع_نهایی_جلسه.pptx'
prs=Presentation(p)
for sn in (6,7,8):
 print('SLIDE',sn)
 for i,sh in enumerate(prs.slides[sn-1].shapes,1):
  print(i,sh.name,sh.shape_type,repr(getattr(sh,'text','').replace('\n',' | ')[:120]))
