from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from zipfile import ZipFile, ZIP_DEFLATED
import os, shutil, tempfile

OUT = r"C:\Users\amir2\Desktop\cat-claude\متن_جایگزین_معماری_YOLOv11.docx"
FA = 'B Zar'
EN = 'Times New Roman'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'

def set_bidi(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1'); pPr.append(bidi)

def set_rtl(run, font, size, bold=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rtl = OxmlElement('w:rtl'); rtl.set(qn('w:val'), '1'); rPr.append(rtl)
    fonts = rPr.rFonts
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rPr.append(fonts)
    fonts.set(qn('w:ascii'), font); fonts.set(qn('w:hAnsi'), font)
    fonts.set(qn('w:cs'), font); fonts.set(qn('w:eastAsia'), font)

def text_run(p, text, latin=False, bold=False):
    r = p.add_run(text)
    set_rtl(r, EN if latin else FA, 12 if latin else 14, bold)
    return r

def ref_run(p, n):
    text_run(p, f' [{n}]')

def footnote_ref(p, fid):
    r = p.add_run()
    rPr = r._element.get_or_add_rPr()
    rStyle = OxmlElement('w:rStyle'); rStyle.set(qn('w:val'), 'Footnote Reference'); rPr.append(rStyle)
    e = OxmlElement('w:footnoteReference'); e.set(qn('w:id'), str(fid)); r._element.append(e)

def add_body(parts):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_bidi(p)
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(2)
    for part in parts:
        if isinstance(part, tuple) and part[0] == 'foot': footnote_ref(p, part[1])
        elif isinstance(part, tuple) and part[0] == 'latin': text_run(p, part[1], latin=True)
        elif isinstance(part, tuple) and part[0] == 'ref': ref_run(p, part[1])
        else: text_run(p, part)
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(3); sec.bottom_margin = Cm(2.5); sec.right_margin = Cm(3); sec.left_margin = Cm(2)
styles = doc.styles
normal = styles['Normal']
normal.font.name = FA; normal.font.size = Pt(14)
normal._element.rPr.rFonts.set(qn('w:cs'), FA)
normal.paragraph_format.line_spacing = 1.5

# Title, kept separate from copyable thesis subsection
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_bidi(p)
p.paragraph_format.space_after = Pt(18)
r = p.add_run('متن جایگزین بخش ۱‏-‏۳‏-‏۴')
set_rtl(r, FA, 16, True)
r = p.add_run(' معماری '); set_rtl(r, FA, 16, True)
r = p.add_run('YOLOv11'); set_rtl(r, EN, 14, True)

# Thesis heading
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_bidi(p)
p.paragraph_format.space_after = Pt(10)
r = p.add_run('۱‏-‏۳‏-‏۴- معماری '); set_rtl(r, FA, 14, True)
r = p.add_run('YOLOv11'); set_rtl(r, EN, 12, True)

add_body([
    'مدل مورد استفاده در این پژوهش، نسخه تشخیص اشیای ', ('latin','YOLOv11'), ' است. معماری این مدل، مانند دیگر مدل‌های جدید خانواده ', ('latin','YOLO'), '، از سه بخش پیوسته تشکیل می‌شود: ستون‌فقرات', ('foot',1), ' برای استخراج نگاشت‌های ویژگی از تصویر ورودی، گردن', ('foot',2), ' برای آمیختن اطلاعات این نگاشت‌ها در مقیاس‌های گوناگون، و سر', ('foot',3), ' برای تولید پیش‌بینی نهایی کادر و دسته شیء. این تقسیم‌بندی، مسیر پردازش تصویر را از استخراج ویژگی‌های اولیه تا پیش‌بینی خروجی مشخص می‌کند', ('ref','۳، ۱۳، ۱۷'), '.'
])
add_body([
    'مطابق پیکربندی رسمی ', ('latin','YOLO11'), '، ستون‌فقرات با پنج لایه پیچشیِ کاهنده آغاز می‌شود و در سه سطح ', ('latin','P3/8'), '، ', ('latin','P4/16'), ' و ', ('latin','P5/32'), ' نگاشت ویژگی تولید می‌کند. عدد پس از خط تیره، نسبت کوچک‌شدن طول و عرض نگاشت ویژگی نسبت به تصویر ورودی را نشان می‌دهد؛ بنابراین شاخه ', ('latin','P3/8'), ' جزئیات مکانی بیشتری دارد و شاخه‌های ', ('latin','P4/16'), ' و ', ('latin','P5/32'), ' زمینه دید وسیع‌تری فراهم می‌کنند. گردن شبکه ابتدا با مسیر بالا به پایین، ویژگی‌های معنایی عمیق را به نگاشت‌های پُرتفکیک‌تر منتقل می‌کند و سپس با مسیر پایین به بالا، اطلاعات مکانی را دوباره به سطوح عمیق‌تر بازمی‌گرداند. سه نگاشت حاصل به سر آشکارساز داده می‌شوند؛ همان سه خروجی‌ای که در شکل ۱‏-‏۲ با برچسب‌های ', ('latin','Detect'), ' نمایش یافته‌اند', ('ref',17), '.'
])
add_body([
    'نخستین تفاوت ساختاری ', ('latin','YOLOv11'), ' با ', ('latin','YOLOv8'), '، جایگزینی بلوک ', ('latin','C2f'), ' با ', ('latin','C3k2'), ' در ستون‌فقرات و گردن است. ', ('latin','C3k2'), ' گونه‌ای از بلوک‌های اتصال جزئی میان‌مرحله‌ای است که واحد تکرارشونده درون آن، بسته به مقیاس مدل، از ساختار ', ('latin','Bottleneck'), ' یا ', ('latin','C3k'), ' استفاده می‌کند. ازاین‌رو، توضیح این بلوک به‌صورت «جایگزین‌کردن یک پیچش بزرگ با دو پیچش کوچک» دقیق نیست. هدف این جایگزینی، پردازش کارآمدتر ویژگی‌ها و کنترل هزینه محاسباتی، در عین حفظ مسیرهای انتقال ویژگی در ساختار ', ('latin','CSP'), ' است', ('ref',17), '.'
])
add_body([
    'پس از آخرین بلوک ', ('latin','C3k2'), ' در مسیر استخراج ویژگی، بلوک ', ('latin','SPPF'), ' قرار دارد که در این نوشتار «ادغام هرمی مکانی سریع»', ('foot',4), ' نامیده می‌شود. در پیکربندی رسمی ', ('latin','YOLO11'), '، این بلوک پیش از نخستین عمل بزرگ‌نمایی در مسیر ادغام چندمقیاسی تعریف شده است؛ بنابراین در مرز انتقال ویژگی‌ها از ستون‌فقرات به گردن قرار می‌گیرد. ', ('latin','SPPF'), ' یک عمل بیشینه‌گیری ', ('latin','۵×۵'), ' را سه بار پیاپی اجرا و خروجی‌های میانی را با ورودی اولیه ترکیب می‌کند. حاصل این کار از نظر میدان دید معادل گردآوری ویژگی‌ها در چند مقیاس است، اما با محاسبه‌ای کم‌هزینه‌تر از اجرای مستقل چند هسته بزرگ. این بلوک در ', ('latin','YOLOv8'), ' نیز وجود دارد و جزء تازه‌ای در ', ('latin','YOLOv11'), ' نیست', ('ref',17), '.'
])
add_body([
    'تغییر دوم، افزوده‌شدن بلوک ', ('latin','C2PSA'), ' بلافاصله پس از ', ('latin','SPPF'), ' است. این بلوک نیز پیش از نخستین عمل بزرگ‌نمایی قرار دارد و خروجی آن وارد مسیر ادغام چندمقیاسی گردن می‌شود. ', ('latin','C2PSA'), ' بخشی از کانال‌های ورودی را از یک مسیر توجه مکانی', ('foot',5), ' عبور می‌دهد و سپس آن را با مسیر میان‌بُر ترکیب می‌کند. درون مسیر توجه، سازوکار توجه چندسری و یک شبکه پیش‌خور کانولوشنی به‌کار می‌روند تا وابستگی میان نواحی تصویر در نگاشت ویژگی مدل‌سازی شود. این سازوکار در ', ('latin','YOLOv8'), ' وجود ندارد؛ بااین‌حال، در این پژوهش آزمایش حذف بلوک انجام نشده است و نمی‌توان سهم مستقل آن را در نتیجه نهایی مدل اندازه‌گیری‌شده دانست', ('ref',17), '.'
])
add_body([
    'سومین تفاوت در سر آشکارساز رخ می‌دهد: شاخه طبقه‌بندی ', ('latin','YOLOv11'), ' با پیچش‌های تفکیک‌پذیر عمقی سبک‌تر شده است. با وجود این تغییر، منطق خروجی نسبت به ', ('latin','YOLOv8'), ' حفظ شده است؛ سر شبکه بدون لنگر است و دو مسیر جداگانه برای برآورد مکان کادر و طبقه‌بندی دارد. مکان کادر با روش ', ('latin','DFL'), ' مدل‌سازی می‌شود و سه خروجی چندمقیاسی برای آشکارسازی اشیای کوچک، متوسط و بزرگ به کار می‌روند. در نتیجه، تفاوت‌های ', ('latin','YOLOv11'), ' با ', ('latin','YOLOv8'), ' تنها به ستون‌فقرات محدود نیست و شامل بلوک‌های ویژگی، توجه و شاخه طبقه‌بندی سر نیز می‌شود', ('ref',17), '.'
])
add_body([
    'این معماری در پنج مقیاس ', ('latin','n'), '، ', ('latin','s'), '، ', ('latin','m'), '، ', ('latin','l'), ' و ', ('latin','x'), ' عرضه می‌شود که با تغییر ضرایب عمق و عرض شبکه، میان شمار پارامترها، بار محاسباتی و دقت مبادله ایجاد می‌کنند. در این پروژه، مقیاس کوچک، یعنی ', ('latin','YOLOv11s'), '، انتخاب شد؛ زیرا نسخه آموزش‌دیده آن برای شش دسته مورد مطالعه حدود ۹٫۴ میلیون پارامتر دارد و با حافظه پردازنده گرافیکی در دسترس سازگار بود. دلیل انتخاب مقیاس کوچک، محدودیت سخت‌افزاری پروژه است؛ این انتخاب به معنای برتری ذاتی این مقیاس در همه کاربردها نیست', ('ref',17), '.'
])

# make all paragraphs keep lines together normally? 
doc.save(OUT)

# Add real footnotes through direct OOXML.
notes = {
 1: 'Backbone', 2: 'Neck', 3: 'Head',
 4: 'Spatial Pyramid Pooling - Fast', 5: 'Spatial Attention'
}
FA_DIGITS = '۱۲۳۴۵'
with tempfile.TemporaryDirectory() as td:
    with ZipFile(OUT) as z: z.extractall(td)
    doc_xml = os.path.join(td,'word','document.xml')
    tree = etree.parse(doc_xml)
    root = tree.getroot(); ns={'w':W}
    # ensure styles named Footnote Text and Footnote Reference
    styles_path=os.path.join(td,'word','styles.xml')
    st=etree.parse(styles_path); sr=st.getroot()
    for sid, typ, name, based in [('FootnoteText','paragraph','footnote text',None),('FootnoteReference','character','footnote reference',None)]:
        if not sr.xpath(f"//w:style[@w:styleId='{sid}']",namespaces=ns):
            e=etree.SubElement(sr,'{%s}style'%W,attrib={'{%s}type'%W:typ,'{%s}styleId'%W:sid})
            etree.SubElement(e,'{%s}name'%W,attrib={'{%s}val'%W:name})
            if based: etree.SubElement(e,'{%s}basedOn'%W,attrib={'{%s}val'%W:based})
            rpr=etree.SubElement(e,'{%s}rPr'%W)
            if sid == 'FootnoteText':
                etree.SubElement(rpr,'{%s}sz'%W,attrib={'{%s}val'%W:'20'})
                etree.SubElement(rpr,'{%s}szCs'%W,attrib={'{%s}val'%W:'20'})
            else:
                etree.SubElement(rpr,'{%s}vertAlign'%W,attrib={'{%s}val'%W:'superscript'})
    st.write(styles_path, xml_declaration=True, encoding='UTF-8', standalone=True)
    foot=etree.Element('{%s}footnotes'%W,nsmap={'w':W})
    for fid, sep in [(-1, 'separator'), (0, 'continuationSeparator')]:
        fn=etree.SubElement(foot,'{%s}footnote'%W,attrib={'{%s}id'%W:str(fid)})
        p=etree.SubElement(fn,'{%s}p'%W); r=etree.SubElement(p,'{%s}r'%W); etree.SubElement(r,'{%s}%s'%(W,sep))
    for fid, txt in notes.items():
        fn=etree.SubElement(foot,'{%s}footnote'%W,attrib={'{%s}id'%W:str(fid)})
        p=etree.SubElement(fn,'{%s}p'%W)
        ppr=etree.SubElement(p,'{%s}pPr'%W); etree.SubElement(ppr,'{%s}pStyle'%W,attrib={'{%s}val'%W:'FootnoteText'})
        r=etree.SubElement(p,'{%s}r'%W); rpr=etree.SubElement(r,'{%s}rPr'%W); etree.SubElement(rpr,'{%s}rFonts'%W,attrib={'{%s}ascii'%W:FA,'{%s}hAnsi'%W:FA,'{%s}cs'%W:FA}); etree.SubElement(rpr,'{%s}sz'%W,attrib={'{%s}val'%W:'20'}); etree.SubElement(rpr,'{%s}szCs'%W,attrib={'{%s}val'%W:'20'}); etree.SubElement(rpr,'{%s}rtl'%W); etree.SubElement(rpr,'{%s}lang'%W,attrib={'{%s}val'%W:'fa-IR','{%s}bidi'%W:'fa-IR'}); tnum=etree.SubElement(r,'{%s}t'%W); tnum.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); tnum.text=FA_DIGITS[fid-1]
        r=etree.SubElement(p,'{%s}r'%W); rpr=etree.SubElement(r,'{%s}rPr'%W); etree.SubElement(rpr,'{%s}rFonts'%W,attrib={'{%s}ascii'%W:EN,'{%s}hAnsi'%W:EN,'{%s}cs'%W:EN}); etree.SubElement(rpr,'{%s}sz'%W,attrib={'{%s}val'%W:'20'}); etree.SubElement(rpr,'{%s}szCs'%W,attrib={'{%s}val'%W:'20'}); etree.SubElement(r,'{%s}t'%W).text=' '+txt
    etree.ElementTree(foot).write(os.path.join(td,'word','footnotes.xml'),xml_declaration=True,encoding='UTF-8',standalone=True)
    # relationship
    relp=os.path.join(td,'word','_rels','document.xml.rels'); rel=etree.parse(relp); rr=rel.getroot()
    if not rr.xpath("//*[local-name()='Relationship' and @Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes']"):
        used=[int(x.get('Id')[3:]) for x in rr if x.get('Id','').startswith('rId') and x.get('Id')[3:].isdigit()]
        etree.SubElement(rr,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='rId'+str(max(used,default=0)+1),Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes',Target='footnotes.xml')
    rel.write(relp,xml_declaration=True,encoding='UTF-8',standalone=True)
    # Content type
    ctp=os.path.join(td,'[Content_Types].xml'); ct=etree.parse(ctp); cr=ct.getroot()
    if not cr.xpath("//*[local-name()='Override' and @PartName='/word/footnotes.xml']"):
        etree.SubElement(cr,'{%s}Override'%CT,PartName='/word/footnotes.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    ct.write(ctp,xml_declaration=True,encoding='UTF-8',standalone=True)
    fixed=OUT+'.tmp'
    with ZipFile(fixed,'w',ZIP_DEFLATED) as z:
        for base,ds,fs in os.walk(td):
            for f in fs:
                full=os.path.join(base,f); z.write(full,os.path.relpath(full,td))
    shutil.move(fixed,OUT)
print(OUT)




