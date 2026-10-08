from pathlib import Path
from docx import Document


src = Path(r"D:\projects\car-detection-yolo\thesis\ارائه\کد\content_a.py")
text = src.read_text(encoding="utf-8")
old = 'sl = slide("انتخاب داده بومی برای تنظیم دقیق مدل")'
new = 'sl = slide("تنظیم دقیق مدل با IADD")'
if text.count(old) != 1:
    raise RuntimeError("slide 6 source title not found exactly once")
src.write_text(text.replace(old, new, 1), encoding="utf-8", newline="\n")

doc_path = Path(r"D:\projects\car-detection-yolo\thesis\ارائه\متن_ارائه.docx")
tmp = doc_path.with_name("متن_ارائه.tmp.docx")
doc = Document(doc_path)
matches = [p for p in doc.paragraphs
           if p.text == "اسلاید ۶ — انتخاب داده بومی برای تنظیم دقیق مدل"]
if len(matches) != 1:
    raise RuntimeError("slide 6 speaker title not found exactly once")
p = matches[0]
p.runs[0].text = "اسلاید ۶ — تنظیم دقیق مدل با IADD"
for run in p.runs[1:]:
    run._element.getparent().remove(run._element)
doc.save(tmp)
Document(tmp)
tmp.replace(doc_path)
print("slide 6 title shortened")
