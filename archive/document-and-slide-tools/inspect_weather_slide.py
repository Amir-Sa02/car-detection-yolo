from pathlib import Path
from pptx import Presentation
p=Path(r'C:\Users\amir2\Desktop\نسخه ارسالی\documents\ارائه_دفاع_نهایی_جلسه.pptx')
prs=Presentation(p)
for i,slide in enumerate(prs.slides,1):
    texts=[s.text.replace('\n',' / ') for s in slide.shapes if s.has_text_frame and s.text.strip()]
    pics=[s for s in slide.shapes if s.shape_type==13]
    print(i,'pics',len(pics),' :: ',' | '.join(texts)[:350])
    if i==27:
        print('slide size', prs.slide_width/914400,prs.slide_height/914400)
        for j,s in enumerate(slide.shapes):
            print(j,s.shape_type,s.name,'XYWH',tuple(round(v/914400,3) for v in (s.left,s.top,s.width,s.height)),
                  'text=',s.text[:90].replace('\n','/') if s.has_text_frame else '',
                  'image=',(s.image.ext,len(s.image.blob)) if s.shape_type==13 else '')
            if s.shape_type==13:
                Path('slide27_'+str(j)+'.'+s.image.ext).write_bytes(s.image.blob)
